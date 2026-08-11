from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TENDER_ID = "Tender_01_VacuumFlask"

LABEL_ALIASES = [
    "Document required from seller",
    "Documents required from seller",
    "Documents required from bidder",
    "Documents asked from seller",
    "विक्रेता से मांगे गए दस्तावेज",
    "विक्रेता से मांगे गये दस्तावेज",
]

STATUS_REVIEW_PENDING = "Pending Review"
STATUS_REVIEW_NOT_REQUIRED = "Not Required"


ATTRIBUTE_HEADERS = [
    "tender_id",
    "requirement_id",
    "requirement_type",
    "requirement_text",
    "expected_bidder_evidence",
    "source_document_id",
    "source_file",
    "source_page",
    "source_label",
    "raw_field_text",
    "raw_item_text",
    "normalized_item_text",
    "canonical_doc_type",
    "requirement_category",
    "evidence_expected",
    "matching_keywords",
    "presence_check_ready",
    "needs_clause_expansion",
    "applicability_condition",
    "is_atc_reference",
    "atc_expansion_status",
    "expected_bidder_doc_group",
    "suggested_filename_keywords",
    "possible_satisfying_documents",
    "strict_document_name_required",
    "human_review_required_for_presence",
    "source_text_snippet",
    "source_snippet",
    "extraction_method",
    "confidence",
    "extraction_confidence",
    "review_status",
    "notes",
]

REQUIREMENT_PREVIEW_HEADERS = [
    "requirement_id",
    "requirement_type",
    "requirement_text",
    "expected_bidder_evidence",
    "source_file",
    "source_page",
    "source_snippet",
    "extraction_confidence",
    "review_status",
]

FIELD_HEADERS = [
    "tender_id",
    "source_document_id",
    "source_file",
    "source_page",
    "source_label",
    "raw_field_text",
    "extraction_method",
    "confidence",
    "notes",
]

TEXT_INDEX_HEADERS = [
    "tender_id",
    "document_id",
    "document_side",
    "bidder_id",
    "bidder_name",
    "file_name",
    "folder_path",
    "page_number",
    "chosen_text_source",
    "native_text_chars",
    "ocr_text_chars",
    "chosen_text_chars",
    "is_scanned_likely",
    "ocr_engine",
    "processing_status",
    "chosen_text",
]


def resolve_tender_root(tender_id: str, tender_root: str | Path | None = None) -> Path:
    return Path(tender_root).resolve() if tender_root else ROOT / tender_id


def output_paths(tender_id: str, tender_root: str | Path | None = None) -> dict[str, Path]:
    tender_root = resolve_tender_root(tender_id, tender_root)
    output_root = tender_root / "05_Extraction_Output"
    text_root = output_root / "01_text_extraction"
    req_root = output_root / "05_tender_requirements"
    return {
        "tender_root": tender_root,
        "output_root": output_root,
        "manifest": output_root / "document_manifest.json",
        "all_pages": text_root / "page_json" / "all_pages.json",
        "text_index_xlsx": text_root / "all_document_text_index.xlsx",
        "text_index_json": text_root / "all_document_text_index.json",
        "req_root": req_root,
        "field_xlsx": req_root / "document_required_from_seller.xlsx",
        "field_json": req_root / "document_required_from_seller.json",
        "attributes_xlsx": req_root / "required_document_attributes.xlsx",
        "attributes_json": req_root / "required_document_attributes.json",
        "bidder_requirements_xlsx": req_root / "bidder_requirements.xlsx",
        "bidder_requirements_json": req_root / "bidder_requirements.json",
        "notes_json": req_root / "extraction_notes.json",
    }


def ensure_dirs(paths: dict[str, Path]) -> None:
    paths["req_root"].mkdir(parents=True, exist_ok=True)
    paths["text_index_xlsx"].parent.mkdir(parents=True, exist_ok=True)


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, rows: Any) -> None:
    path.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")


def autosize(ws) -> None:
    for column_cells in ws.columns:
        col = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col].width = min(max(max_len + 2, 12), 70)


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append([excel_safe(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)
    wb.save(path)


def excel_safe(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", value)
    return value


def normalize_for_match(text: str) -> str:
    text = text.casefold()
    text = re.sub(r"[\u200b-\u200f\ufeff]", "", text)
    text = re.sub(r"[^0-9a-z\u0900-\u097f]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def compact_for_match(text: str) -> str:
    return re.sub(r"\s+", "", normalize_for_match(text))


def manifest_by_document_id(manifest_rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {row["document_id"]: row for row in manifest_rows}


def build_text_index(
    tender_id: str,
    manifest_rows: list[dict[str, Any]],
    page_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    manifest_lookup = manifest_by_document_id(manifest_rows)
    rows: list[dict[str, Any]] = []
    for page in page_rows:
        doc = manifest_lookup.get(page["document_id"], {})
        chosen_text = page.get("chosen_text", "") or ""
        rows.append(
            {
                "tender_id": tender_id,
                "document_id": page.get("document_id", ""),
                "document_side": doc.get("document_side", ""),
                "bidder_id": doc.get("bidder_id", ""),
                "bidder_name": doc.get("bidder_name", ""),
                "file_name": page.get("file_name", doc.get("file_name", "")),
                "folder_path": doc.get("folder_path", ""),
                "page_number": page.get("page_number", ""),
                "chosen_text_source": page.get("chosen_text_source", ""),
                "native_text_chars": page.get("native_text_chars", 0),
                "ocr_text_chars": page.get("ocr_text_chars", 0),
                "chosen_text_chars": len(chosen_text.strip()),
                "is_scanned_likely": page.get("is_scanned_likely", ""),
                "ocr_engine": page.get("ocr_engine", ""),
                "processing_status": page.get("processing_status", ""),
                "chosen_text": chosen_text,
            }
        )
    return rows


def label_patterns() -> list[tuple[str, re.Pattern[str]]]:
    patterns: list[tuple[str, re.Pattern[str]]] = []
    english_core = [
        r"documents?\s+required\s+from\s+(?:seller|bidder)",
        r"documents?\s+asked\s+from\s+seller",
        r"document\s+required\s+from\s+seller",
    ]
    for pattern in english_core:
        patterns.append((pattern, re.compile(pattern, re.IGNORECASE)))
    hindi_core = [
        r"विक्रेता\s*से\s*मांगे\s*गए\s*दस्तावेज",
        r"विक्रेता\s*से\s*मांगे\s*गये\s*दस्तावेज",
    ]
    for pattern in hindi_core:
        patterns.append((pattern, re.compile(pattern)))
    return patterns


def find_label(text: str) -> tuple[str, int, int] | None:
    for label, pattern in label_patterns():
        match = pattern.search(text)
        if match:
            return label, match.start(), match.end()
    compact_text = compact_for_match(text)
    for alias in LABEL_ALIASES:
        compact_alias = compact_for_match(alias)
        idx = compact_text.find(compact_alias)
        if idx >= 0:
            # Compact matching cannot reliably map positions, so use a fallback English location.
            lower = text.casefold()
            for token in ("document", "विक्रेता"):
                raw_idx = lower.find(token)
                if raw_idx >= 0:
                    return alias, raw_idx, raw_idx + len(token)
            return alias, 0, 0
    return None


def end_boundary(text_after_label: str) -> int:
    boundary_patterns = [
        r"\n\s*/?\s*Bid\s+Number\s*:",
        r"\n\s*Bid\s+Number\s*:",
        r"\n\s*/?\s*Do\s+you\s+want\s+to\s+show",
        r"\n\s*---\s*Page\s+\d+\s*---",
        r"\n\s*\d+\s*/\s*\d+\s*$",
    ]
    candidates = []
    for pattern in boundary_patterns:
        match = re.search(pattern, text_after_label, flags=re.IGNORECASE)
        if match:
            candidates.append(match.start())
    return min(candidates) if candidates else min(len(text_after_label), 2000)


def clean_field_text(text: str) -> str:
    text = re.sub(r"[\x00-\x08\x0b-\x1f]+", " ", text)
    text = text.replace("\r", "\n")
    lines = [line.strip(" /|\t") for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines).strip()


def extract_field_blocks(
    tender_id: str,
    manifest_rows: list[dict[str, Any]],
    page_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    manifest_lookup = manifest_by_document_id(manifest_rows)
    blocks: list[dict[str, Any]] = []
    for page in page_rows:
        doc = manifest_lookup.get(page["document_id"], {})
        if doc.get("document_side") != "tender-side":
            continue
        text = page.get("chosen_text", "") or ""
        label_hit = find_label(text)
        if not label_hit:
            continue
        label, start, end = label_hit
        after = text[end:]
        boundary = end_boundary(after)
        raw_field = clean_field_text(after[:boundary])
        # The Hindi label sometimes appears before the English label. Trim duplicate label fragments.
        raw_field = re.sub(r"^from\s+seller\s*", "", raw_field, flags=re.IGNORECASE).strip()
        raw_field = re.sub(r"^required\s+from\s+(?:seller|bidder)\s*", "", raw_field, flags=re.IGNORECASE).strip()
        blocks.append(
            {
                "tender_id": tender_id,
                "source_document_id": page["document_id"],
                "source_file": page.get("file_name", doc.get("file_name", "")),
                "source_page": page.get("page_number", ""),
                "source_label": label,
                "raw_field_text": raw_field,
                "extraction_method": "label_block_parser",
                "confidence": 0.9 if "Document required" in text[start : start + 80] else 0.75,
                "notes": "",
                "_snippet": source_snippet(text, start, end + boundary),
            }
        )
    return blocks


def source_snippet(text: str, start: int, end: int, width: int = 220) -> str:
    left = max(0, start - width)
    right = min(len(text), end + width)
    return re.sub(r"\s+", " ", text[left:right]).strip()


def split_preserving_parentheses(text: str) -> list[str]:
    item_block = re.split(r"\*?\s*In\s+case\s+any\s+bidder", text, maxsplit=1, flags=re.IGNORECASE)[0]
    item_block = re.split(r"\bBid\s+Number\s*:", item_block, maxsplit=1, flags=re.IGNORECASE)[0]
    if "," in item_block or ";" in item_block:
        item_block = re.sub(r"\s*\n\s*", " ", item_block)
    else:
        item_block = re.sub(r"\n\s*(?=[*•●▪]|\d+[\).])", "\n", item_block)
    text = item_block
    protected = []

    def repl(match: re.Match[str]) -> str:
        protected.append(match.group(0))
        return f"__PAREN_{len(protected) - 1}__"

    working = re.sub(r"\([^)]*\)", repl, text)
    working = re.sub(r"[*•●▪]", "\n", working)
    working = re.sub(r"(?m)^\s*\d+[\).]\s*", "\n", working)
    parts: list[str] = []
    current = []
    for char in working:
        if char in ",;\n":
            item = "".join(current).strip()
            if item:
                parts.append(item)
            current = []
        else:
            current.append(char)
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    restored = []
    for part in parts:
        for idx, value in enumerate(protected):
            part = part.replace(f"__PAREN_{idx}__", value)
        item = clean_item(part)
        if item:
            restored.append(item)
    return restored


def clean_item(item: str) -> str:
    item = re.sub(r"\s+", " ", item).strip(" .;:-")
    item = item.strip()
    if not item:
        return ""
    if item.lower().startswith("in case any bidder") or item.startswith("*In case any bidder"):
        return ""
    if "supporting documents to prove" in item.lower():
        return ""
    if len(item) > 120:
        return ""
    return item


def normalized_item_text(item: str) -> str:
    item = re.sub(r"\s+", " ", item).strip()
    return item[:1].upper() + item[1:] if item else item


def preview_type_from_category(category: str, item: str = "") -> str:
    norm = normalize_for_match(item)
    if category == "financial_turnover" or "turnover" in norm:
        return "financial_criteria"
    if category in {"experience", "past_performance"}:
        return "experience_criteria"
    if category in {"oem_authorization", "oem_turnover"} or "oem" in norm:
        return "oem_condition"
    if category == "atc_reference" or "atc" in norm:
        return "atc_reference"
    if "declaration" in norm or "undertaking" in norm:
        return "declaration"
    return "document"


def mapping_for_item(item: str) -> dict[str, Any]:
    norm = normalize_for_match(item)
    if "certificate requested in atc" in norm:
        return {
            "canonical_doc_type": "ATC requested certificate",
            "requirement_category": "atc_reference",
            "evidence_expected": "Certificate or document requested in ATC clauses",
            "matching_keywords": "certificate, requested in atc, atc",
            "presence_check_ready": "No",
            "needs_clause_expansion": "Yes",
            "applicability_condition": "Requires ATC clause expansion before presence check.",
            "is_atc_reference": "Yes",
            "atc_expansion_status": "Deferred",
            "expected_bidder_doc_group": "ATC referenced documents",
            "suggested_filename_keywords": "certificate, atc, annexure",
            "possible_satisfying_documents": "Depends on ATC requirement; expansion deferred",
            "strict_document_name_required": "No",
            "human_review_required_for_presence": "Yes",
            "review_status": STATUS_REVIEW_PENDING,
            "notes": "ATC reference captured; not expanded in this step.",
        }
    if "experience criteria" in norm:
        return {
            "canonical_doc_type": "Experience / CRAC / work order evidence",
            "requirement_category": "experience",
            "evidence_expected": "CRAC, work order, completion certificate, material receipt certificate, or similar experience evidence",
            "matching_keywords": "experience, crac, work order, completion, material receipt",
            "presence_check_ready": "Yes",
            "needs_clause_expansion": "No",
            "applicability_condition": "Required unless tender exemption applies.",
            "is_atc_reference": "No",
            "atc_expansion_status": "Not Applicable",
            "expected_bidder_doc_group": "Experience evidence",
            "suggested_filename_keywords": "experience, crac, work, order, completion, receipt",
            "possible_satisfying_documents": "CRAC, work order, completion certificate, material receipt certificate",
            "strict_document_name_required": "No",
            "human_review_required_for_presence": "No",
            "review_status": STATUS_REVIEW_NOT_REQUIRED,
            "notes": "",
        }
    if "past performance" in norm:
        return {
            "canonical_doc_type": "Past performance evidence",
            "requirement_category": "past_performance",
            "evidence_expected": "Past performance or prior supply evidence",
            "matching_keywords": "past performance, performance, past supply, supplied",
            "presence_check_ready": "Yes",
            "needs_clause_expansion": "No",
            "applicability_condition": "Required unless tender exemption applies.",
            "is_atc_reference": "No",
            "atc_expansion_status": "Not Applicable",
            "expected_bidder_doc_group": "Past performance evidence",
            "suggested_filename_keywords": "past, performance, supply, supplied",
            "possible_satisfying_documents": "Past performance certificate, supply proof, GeM performance document",
            "strict_document_name_required": "No",
            "human_review_required_for_presence": "No",
            "review_status": STATUS_REVIEW_NOT_REQUIRED,
            "notes": "",
        }
    if "bidder turnover" in norm:
        return {
            "canonical_doc_type": "Bidder turnover evidence",
            "requirement_category": "financial_turnover",
            "evidence_expected": "Audited financial statements or bidder turnover proof",
            "matching_keywords": "bidder turnover, turnover, audited, balance sheet, profit, loss, ca",
            "presence_check_ready": "Yes",
            "needs_clause_expansion": "No",
            "applicability_condition": "Required unless tender exemption applies.",
            "is_atc_reference": "No",
            "atc_expansion_status": "Not Applicable",
            "expected_bidder_doc_group": "Bidder financial evidence",
            "suggested_filename_keywords": "turnover, audited, balance, financial, ca",
            "possible_satisfying_documents": "Audited balance sheet, P&L statement, turnover certificate where accepted",
            "strict_document_name_required": "No",
            "human_review_required_for_presence": "No",
            "review_status": STATUS_REVIEW_NOT_REQUIRED,
            "notes": "",
        }
    if "oem authorization certificate" in norm or "oem authorisation certificate" in norm:
        return {
            "canonical_doc_type": "OEM authorization certificate",
            "requirement_category": "oem_authorization",
            "evidence_expected": "OEM authorization certificate",
            "matching_keywords": "oem, authorization, authorisation, certificate",
            "presence_check_ready": "Yes",
            "needs_clause_expansion": "No",
            "applicability_condition": "Required for offered product where OEM authorization is applicable.",
            "is_atc_reference": "No",
            "atc_expansion_status": "Not Applicable",
            "expected_bidder_doc_group": "OEM authorization",
            "suggested_filename_keywords": "oem, authorization, authorisation",
            "possible_satisfying_documents": "OEM authorization certificate, manufacturer authorization form",
            "strict_document_name_required": "No",
            "human_review_required_for_presence": "No",
            "review_status": STATUS_REVIEW_NOT_REQUIRED,
            "notes": "",
        }
    if "oem annual turnover" in norm or ("oem" in norm and "turnover" in norm):
        return {
            "canonical_doc_type": "OEM turnover evidence",
            "requirement_category": "oem_turnover",
            "evidence_expected": "OEM annual turnover proof",
            "matching_keywords": "oem, annual turnover, turnover, ca, audited",
            "presence_check_ready": "Yes",
            "needs_clause_expansion": "No",
            "applicability_condition": "Required for OEM turnover criteria where applicable.",
            "is_atc_reference": "No",
            "atc_expansion_status": "Not Applicable",
            "expected_bidder_doc_group": "OEM financial evidence",
            "suggested_filename_keywords": "oem, turnover, annual, ca, audited",
            "possible_satisfying_documents": "OEM turnover certificate, OEM audited financial statement",
            "strict_document_name_required": "No",
            "human_review_required_for_presence": "No",
            "review_status": STATUS_REVIEW_NOT_REQUIRED,
            "notes": "",
        }
    return {
        "canonical_doc_type": "Unmapped required document",
        "requirement_category": "unknown",
        "evidence_expected": item,
        "matching_keywords": item,
        "presence_check_ready": "Yes",
        "needs_clause_expansion": "No",
        "applicability_condition": "Explicit tender document item; keyword-based presence check is allowed.",
        "is_atc_reference": "No",
        "atc_expansion_status": "Not Applicable",
        "expected_bidder_doc_group": "Tender requested document",
        "suggested_filename_keywords": item,
        "possible_satisfying_documents": item,
        "strict_document_name_required": "No",
        "human_review_required_for_presence": "No",
        "review_status": STATUS_REVIEW_PENDING,
        "notes": "Unmapped document item retained and checked by filename/type keywords.",
    }


def build_attributes(tender_id: str, field_blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    counter = 1
    for block in field_blocks:
        items = split_preserving_parentheses(block["raw_field_text"])
        for item in items:
            mapping = mapping_for_item(item)
            requirement_type = preview_type_from_category(mapping["requirement_category"], item)
            requirement_text = normalized_item_text(item)
            row = {
                "tender_id": tender_id,
                "requirement_id": f"REQ-DOC-{counter:03d}",
                "requirement_type": requirement_type,
                "requirement_text": requirement_text,
                "expected_bidder_evidence": mapping["evidence_expected"],
                "source_document_id": block["source_document_id"],
                "source_file": block["source_file"],
                "source_page": block["source_page"],
                "source_label": block["source_label"],
                "raw_field_text": block["raw_field_text"],
                "raw_item_text": item,
                "normalized_item_text": requirement_text,
                "source_text_snippet": block["_snippet"],
                "source_snippet": block["_snippet"],
                "extraction_method": block["extraction_method"],
                "confidence": block["confidence"],
                "extraction_confidence": block["confidence"],
                **mapping,
            }
            rows.append(row)
            counter += 1
    return rows


CONDITION_RULES: list[dict[str, Any]] = [
    {
        "requirement_type": "financial_criteria",
        "canonical_doc_type": "Bidder turnover criteria",
        "requirement_category": "financial_turnover",
        "evidence_expected": "Bidder turnover proof or audited financial evidence",
        "pattern": r"minimum\s+average\s+annual\s+turnover\s+of\s+the\s+bidder|bidder\s+turnover|financial\s+turnover",
    },
    {
        "requirement_type": "financial_criteria",
        "canonical_doc_type": "OEM turnover criteria",
        "requirement_category": "oem_turnover",
        "evidence_expected": "OEM turnover proof or audited financial evidence",
        "pattern": r"oem\s+average\s+turnover|oem\s+annual\s+turnover|oem.*turnover",
    },
    {
        "requirement_type": "experience_criteria",
        "canonical_doc_type": "Experience criteria",
        "requirement_category": "experience",
        "evidence_expected": "Experience certificate, work order, completion certificate, CRAC, or similar evidence",
        "pattern": r"experience\s+criteria|past\s+experience|similar\s+work|work\s+order|completion\s+certificate|crac|past\s+performance",
    },
    {
        "requirement_type": "oem_condition",
        "canonical_doc_type": "OEM authorization condition",
        "requirement_category": "oem_authorization",
        "evidence_expected": "OEM authorization certificate or manufacturer authorization evidence",
        "pattern": r"oem\s+authori[sz]ation|manufacturer\s+authori[sz]ation|authori[sz]ation\s+certificate",
    },
    {
        "requirement_type": "technical_condition",
        "canonical_doc_type": "Technical criteria",
        "requirement_category": "technical_condition",
        "evidence_expected": "Technical compliance document, specification sheet, catalogue, or supporting certificate",
        "pattern": r"technical\s+(?:specification|criteria|compliance)|compliance\s+sheet|specification\s+document",
    },
    {
        "requirement_type": "declaration",
        "canonical_doc_type": "Declaration / undertaking",
        "requirement_category": "declaration",
        "evidence_expected": "Declaration, undertaking, certificate, affidavit, or annexure requested by the tender",
        "pattern": r"declaration|undertaking|self\s+certification|certificate\s+to\s+be\s+submitted|annexure",
    },
    {
        "requirement_type": "atc_reference",
        "canonical_doc_type": "ATC referenced condition",
        "requirement_category": "atc_reference",
        "evidence_expected": "Document or evidence requested in Buyer Added Bid Specific ATC",
        "pattern": r"buyer\s+added\s+bid\s+specific\s+atc|requested\s+in\s+atc|certificate\s*\(requested\s+in\s+atc\)|atc\s+clause",
    },
    {
        "requirement_type": "other_condition",
        "canonical_doc_type": "Applicability / exemption condition",
        "requirement_category": "other_condition",
        "evidence_expected": "Applicability or exemption evidence, where relevant",
        "pattern": r"\bmse\b|startup|start-up|exemption|exempted|applicable\s+for|not\s+applicable",
    },
]


def condition_candidate_text(text: str, start: int, end: int) -> str:
    line_start = text.rfind("\n", 0, start)
    line_end = text.find("\n", end)
    if line_start < 0:
        line_start = max(0, start - 220)
    if line_end < 0:
        line_end = min(len(text), end + 220)
    left = max(0, line_start)
    right = min(len(text), line_end)
    candidate = clean_item(clean_field_text(text[left:right]))
    if candidate and 18 <= len(candidate) <= 220:
        return candidate
    return source_snippet(text, start, end, width=140)


def build_condition_rows(
    tender_id: str,
    manifest_rows: list[dict[str, Any]],
    page_rows: list[dict[str, Any]],
    existing_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    manifest_lookup = manifest_by_document_id(manifest_rows)
    seen = {
        (row.get("requirement_type", ""), normalize_for_match(row.get("requirement_text") or row.get("normalized_item_text", "")))
        for row in existing_rows
    }
    rows: list[dict[str, Any]] = []
    counter = 1
    for page in page_rows:
        doc = manifest_lookup.get(page["document_id"], {})
        if doc.get("document_side") != "tender-side":
            continue
        text = page.get("chosen_text", "") or ""
        for rule in CONDITION_RULES:
            for match in re.finditer(rule["pattern"], text, flags=re.IGNORECASE):
                requirement_text = normalized_item_text(condition_candidate_text(text, match.start(), match.end()))
                normalized_key = normalize_for_match(requirement_text)
                if not normalized_key:
                    continue
                key = (rule["requirement_type"], normalized_key)
                if key in seen:
                    continue
                seen.add(key)
                snippet = source_snippet(text, match.start(), match.end())
                rows.append(
                    {
                        "tender_id": tender_id,
                        "requirement_id": f"REQ-CON-{counter:03d}",
                        "requirement_type": rule["requirement_type"],
                        "requirement_text": requirement_text,
                        "expected_bidder_evidence": rule["evidence_expected"],
                        "source_document_id": page["document_id"],
                        "source_file": page.get("file_name", doc.get("file_name", "")),
                        "source_page": page.get("page_number", ""),
                        "source_label": rule["canonical_doc_type"],
                        "raw_field_text": requirement_text,
                        "raw_item_text": requirement_text,
                        "normalized_item_text": requirement_text,
                        "canonical_doc_type": rule["canonical_doc_type"],
                        "requirement_category": rule["requirement_category"],
                        "evidence_expected": rule["evidence_expected"],
                        "matching_keywords": requirement_text,
                        "presence_check_ready": "No",
                        "needs_clause_expansion": "Yes",
                        "applicability_condition": "Semantic criteria extracted for reviewer confirmation.",
                        "is_atc_reference": "Yes" if rule["requirement_type"] == "atc_reference" else "No",
                        "atc_expansion_status": "Pending Review" if rule["requirement_type"] == "atc_reference" else "Not Applicable",
                        "expected_bidder_doc_group": rule["canonical_doc_type"],
                        "suggested_filename_keywords": requirement_text,
                        "possible_satisfying_documents": rule["evidence_expected"],
                        "strict_document_name_required": "No",
                        "human_review_required_for_presence": "Yes",
                        "source_text_snippet": snippet,
                        "source_snippet": snippet,
                        "extraction_method": "condition_keyword_parser",
                        "confidence": 0.65,
                        "extraction_confidence": 0.65,
                        "review_status": STATUS_REVIEW_PENDING,
                        "notes": "Criteria or condition extracted from tender text; bidder-wise evaluation requires human review or a dedicated evaluator.",
                    }
                )
                counter += 1
    return rows


def preview_rows(attribute_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for row in attribute_rows:
        rows.append({header: row.get(header, "") for header in REQUIREMENT_PREVIEW_HEADERS})
    return rows


def extraction_notes(
    tender_id: str,
    text_index_rows: list[dict[str, Any]],
    field_blocks: list[dict[str, Any]],
    attribute_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "tender_id": tender_id,
        "text_index_rows": len(text_index_rows),
        "field_blocks_found": len(field_blocks),
        "attribute_rows": len(attribute_rows),
        "bidder_requirement_rows": len(attribute_rows),
        "tender_side_pages_scanned": sum(1 for row in text_index_rows if row["document_side"] == "tender-side"),
        "bidder_side_pages_ignored_for_requirement_extraction": sum(1 for row in text_index_rows if row["document_side"] == "bidder-side"),
        "label_aliases": LABEL_ALIASES,
        "notes": [
            "Requirement extraction uses only tender-side chosen_text.",
            "Explicit documents, criteria, conditions, exemptions, and ATC references are displayed for reviewer confirmation.",
            "Criteria and ATC references are captured as Needs Review unless a dedicated evaluator exists.",
            "Bidder submission evaluation is not performed in this step.",
        ],
    }


def validate_outputs(
    text_index_rows: list[dict[str, Any]],
    field_blocks: list[dict[str, Any]],
    attribute_rows: list[dict[str, Any]],
    all_pages: list[dict[str, Any]],
) -> list[str]:
    issues: list[str] = []
    if len(text_index_rows) != len(all_pages):
        issues.append(f"Text index row count {len(text_index_rows)} does not match page row count {len(all_pages)}.")
    if any("bidder" in str(row.get("document_side", "")).lower() for row in field_blocks):
        issues.append("A bidder-side field block was used unexpectedly.")
    if not attribute_rows:
        issues.append("No bidder requirements extracted from tender-side text.")
    for row in attribute_rows:
        if row["normalized_item_text"].casefold().startswith("in case any bidder"):
            issues.append("Exemption note was emitted as a required-document item.")
    return issues


def run(tender_id: str, tender_root: str | Path | None = None) -> int:
    paths = output_paths(tender_id, tender_root)
    ensure_dirs(paths)
    manifest_rows = load_json(paths["manifest"])
    all_pages = load_json(paths["all_pages"])
    text_index_rows = build_text_index(tender_id, manifest_rows, all_pages)
    field_blocks = extract_field_blocks(tender_id, manifest_rows, all_pages)
    document_rows = build_attributes(tender_id, field_blocks)
    condition_rows = build_condition_rows(tender_id, manifest_rows, all_pages, document_rows)
    attribute_rows = document_rows + condition_rows
    notes = extraction_notes(tender_id, text_index_rows, field_blocks, attribute_rows)
    issues = validate_outputs(text_index_rows, field_blocks, attribute_rows, all_pages)
    notes["validation_issues"] = issues

    write_xlsx(paths["text_index_xlsx"], "All Document Text", text_index_rows, TEXT_INDEX_HEADERS)
    write_json(paths["text_index_json"], text_index_rows)
    write_xlsx(paths["field_xlsx"], "Document Required", field_blocks, FIELD_HEADERS)
    write_json(paths["field_json"], [{k: v for k, v in row.items() if not k.startswith("_")} for row in field_blocks])
    write_xlsx(paths["attributes_xlsx"], "Required Attributes", attribute_rows, ATTRIBUTE_HEADERS)
    write_json(paths["attributes_json"], attribute_rows)
    write_xlsx(paths["bidder_requirements_xlsx"], "Bidder Requirements", preview_rows(attribute_rows), REQUIREMENT_PREVIEW_HEADERS)
    write_json(paths["bidder_requirements_json"], preview_rows(attribute_rows))
    write_json(paths["notes_json"], notes)

    print(f"Text index rows: {len(text_index_rows)}")
    print(f"Field blocks found: {len(field_blocks)}")
    print(f"Bidder requirement rows: {len(attribute_rows)}")
    print(f"Validation issues: {len(issues)}")
    for issue in issues:
        print(f"- {issue}")
    return 1 if issues else 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract GeM tender bidder requirements.")
    parser.add_argument("--tender-id", default=DEFAULT_TENDER_ID)
    parser.add_argument("--tender-root", default="", help="Optional explicit tender workspace root.")
    args = parser.parse_args()
    raise SystemExit(run(args.tender_id, args.tender_root or None))


if __name__ == "__main__":
    main()
