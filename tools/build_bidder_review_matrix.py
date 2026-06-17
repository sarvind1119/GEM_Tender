from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TENDER_ID = "Tender_01_VacuumFlask"

STATUS_PRESENT = "Present"
STATUS_REVIEW = "Needs Review"
FORBIDDEN_TERMS = ("missing", "non-compliant", "rejected", "accepted", "disqualified")

CORE_HEADERS = [
    "tender_id",
    "bidder_id",
    "bidder_name",
    "pan",
    "pan_status",
    "pan_source_file",
    "pan_page",
    "gstin",
    "gstin_status",
    "gstin_source_file",
    "gstin_page",
    "gstin_pan_match_status",
    "identity_needs_human_review",
]

REVIEWER_HEADERS = [
    "reviewer_decision",
    "reviewer_notes",
    "overall_needs_human_review",
]

NOTES_HEADERS = [
    "tender_id",
    "bidder_count",
    "required_document_count",
    "matrix_rows",
    "matrix_columns",
    "json_output",
    "xlsx_output",
    "validation_issues",
]


def resolve_tender_root(tender_id: str, tender_root: str | Path | None = None) -> Path:
    return Path(tender_root).resolve() if tender_root else ROOT / tender_id


def paths(tender_id: str, tender_root: str | Path | None = None) -> dict[str, Path]:
    tender_root = resolve_tender_root(tender_id, tender_root)
    output_root = tender_root / "05_Extraction_Output"
    matrix_root = output_root / "06_bidder_review_matrix"
    return {
        "manifest": output_root / "document_manifest.json",
        "identity": output_root / "02_identity_extraction" / "bidder_identity_summary.json",
        "requirements": output_root / "05_tender_requirements" / "required_document_attributes.json",
        "matrix_root": matrix_root,
        "matrix_xlsx": matrix_root / "bidder_document_review_matrix.xlsx",
        "matrix_json": matrix_root / "bidder_document_review_matrix.json",
        "notes_json": matrix_root / "bidder_document_review_matrix_notes.json",
        "notes_xlsx": matrix_root / "bidder_document_review_matrix_notes.xlsx",
    }


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")


def excel_safe(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", value)
    if isinstance(value, (list, dict)):
        return ILLEGAL_CHARACTERS_RE.sub(" ", json.dumps(value, ensure_ascii=False))
    return value


def autosize(ws) -> None:
    for column_cells in ws.columns:
        col = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col].width = min(max(max_len + 2, 12), 55)


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append([excel_safe(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)
    wb.save(path)


def normalize_header(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    value = value.replace("/", " ")
    return value


def normalize_text(value: str) -> str:
    value = value.casefold()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def split_keywords(*values: str) -> list[str]:
    keywords: list[str] = []
    for value in values:
        for part in re.split(r"[,;|]", value or ""):
            part = normalize_text(part)
            if not part:
                continue
            if part in {"certificate", "document", "evidence", "proof"}:
                continue
            keywords.append(part)
    seen: set[str] = set()
    unique: list[str] = []
    for keyword in keywords:
        if keyword not in seen:
            unique.append(keyword)
            seen.add(keyword)
    return unique


def bidder_docs(manifest_rows: list[dict[str, Any]], bidder_id: str) -> list[dict[str, Any]]:
    return [
        row
        for row in manifest_rows
        if row.get("document_side") == "bidder-side" and row.get("bidder_id") == bidder_id
    ]


def doc_haystack(doc: dict[str, Any]) -> str:
    return normalize_text(
        " ".join(
            [
                str(doc.get("document_type_guess", "")),
                str(doc.get("file_name", "")),
                str(doc.get("folder_path", "")),
            ]
        )
    )


def category_match(requirement: dict[str, Any], doc: dict[str, Any]) -> bool:
    category = requirement.get("requirement_category", "")
    haystack = doc_haystack(doc)

    if category == "experience":
        return any(token in haystack for token in ("experience", "crac", "work order", "work", "completion", "receipt"))
    if category == "past_performance":
        return any(token in haystack for token in ("past", "performance", "experience", "supplied", "supply"))
    if category == "financial_turnover":
        return "turnover" in haystack and "oem" not in haystack
    if category == "oem_authorization":
        return "oem" in haystack and ("authorization" in haystack or "authorisation" in haystack)
    if category == "oem_turnover":
        return "oem" in haystack and "turnover" in haystack
    return False


def keyword_match(requirement: dict[str, Any], doc: dict[str, Any]) -> bool:
    haystack = doc_haystack(doc)
    keywords = split_keywords(
        requirement.get("matching_keywords", ""),
        requirement.get("suggested_filename_keywords", ""),
        requirement.get("canonical_doc_type", ""),
    )
    if not keywords:
        return False
    score = 0
    for keyword in keywords:
        if keyword in haystack:
            score += 2 if " " in keyword else 1
    return score >= 2


def match_requirement_to_docs(requirement: dict[str, Any], docs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if requirement.get("presence_check_ready") != "Yes":
        return []
    matches = [doc for doc in docs if category_match(requirement, doc)]
    if matches:
        return matches
    if requirement.get("requirement_category") not in {"unknown", "", None}:
        return []
    return [doc for doc in docs if keyword_match(requirement, doc)]


def status_for_requirement(requirement: dict[str, Any], matches: list[dict[str, Any]]) -> tuple[str, str]:
    if requirement.get("presence_check_ready") != "Yes":
        return STATUS_REVIEW, "ATC expansion deferred; cannot presence-check until clause expansion is implemented."
    if matches:
        return STATUS_PRESENT, "Likely source document identified from filename/type keywords."
    return STATUS_REVIEW, "No likely source document identified from filename/type keywords."


def requirement_column_groups(requirements: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    for requirement in requirements:
        label = normalize_header(requirement["normalized_item_text"])
        columns.extend(
            [
                f"{label}_status",
                f"{label}_source_files",
                f"{label}_document_ids",
                f"{label}_review_notes",
            ]
        )
    return columns


def build_matrix_rows(
    manifest_rows: list[dict[str, Any]],
    identity_rows: list[dict[str, Any]],
    requirements: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for identity in sorted(identity_rows, key=lambda row: row.get("bidder_id", "")):
        bidder_id = identity.get("bidder_id", "")
        docs = bidder_docs(manifest_rows, bidder_id)
        row: dict[str, Any] = {
            "tender_id": identity.get("tender_id", ""),
            "bidder_id": bidder_id,
            "bidder_name": identity.get("bidder_name", ""),
            "pan": identity.get("pan", identity.get("preferred_pan", "")),
            "pan_status": identity.get("pan_status", identity.get("preferred_pan_status", "")),
            "pan_source_file": identity.get("pan_source_file", identity.get("preferred_pan_source_file", "")),
            "pan_page": identity.get("pan_page", identity.get("preferred_pan_page", "")),
            "gstin": identity.get("gstin", identity.get("preferred_gstin", "")),
            "gstin_status": identity.get("gstin_status", identity.get("preferred_gstin_status", "")),
            "gstin_source_file": identity.get("gstin_source_file", identity.get("preferred_gstin_source_file", "")),
            "gstin_page": identity.get("gstin_page", identity.get("preferred_gstin_page", "")),
            "gstin_pan_match_status": identity.get("gstin_pan_match_status", ""),
            "identity_needs_human_review": identity.get("needs_human_review", ""),
        }
        overall_needs_review = row["identity_needs_human_review"] == "Yes"
        for requirement in requirements:
            label = normalize_header(requirement["normalized_item_text"])
            matches = match_requirement_to_docs(requirement, docs)
            status, note = status_for_requirement(requirement, matches)
            if status == STATUS_REVIEW:
                overall_needs_review = True
            row[f"{label}_status"] = status
            row[f"{label}_source_files"] = "; ".join(doc.get("file_name", "") for doc in matches)
            row[f"{label}_document_ids"] = "; ".join(doc.get("document_id", "") for doc in matches)
            row[f"{label}_review_notes"] = note
        row["reviewer_decision"] = ""
        row["reviewer_notes"] = ""
        row["overall_needs_human_review"] = "Yes" if overall_needs_review else "No"
        rows.append(row)
    return rows


def validate_rows(rows: list[dict[str, Any]], requirements: list[dict[str, Any]], identity_rows: list[dict[str, Any]]) -> list[str]:
    issues: list[str] = []
    if len(rows) != len(identity_rows):
        issues.append(f"Matrix row count {len(rows)} does not match identity row count {len(identity_rows)}.")
    for requirement in requirements:
        label = normalize_header(requirement["normalized_item_text"])
        required_cols = [
            f"{label}_status",
            f"{label}_source_files",
            f"{label}_document_ids",
            f"{label}_review_notes",
        ]
        for col in required_cols:
            if any(col not in row for row in rows):
                issues.append(f"Missing requirement column: {col}")
    for row in rows:
        text = json.dumps(row, ensure_ascii=False).casefold()
        for term in FORBIDDEN_TERMS:
            if term in text:
                issues.append(f"Forbidden term '{term}' appears in row for {row.get('bidder_id')}.")
    return issues


def build_notes(
    tender_id: str,
    rows: list[dict[str, Any]],
    requirements: list[dict[str, Any]],
    headers: list[str],
    output_paths: dict[str, Path],
    validation_issues: list[str],
) -> dict[str, Any]:
    return {
        "tender_id": tender_id,
        "bidder_count": len(rows),
        "required_document_count": len(requirements),
        "matrix_rows": len(rows),
        "matrix_columns": len(headers),
        "json_output": str(output_paths["matrix_json"]),
        "xlsx_output": str(output_paths["matrix_xlsx"]),
        "validation_issues": validation_issues,
        "status_vocabulary": [STATUS_PRESENT, STATUS_REVIEW],
        "notes": [
            "This matrix is a reviewer aid, not a final compliance decision.",
            "Presence matching is based on document type and filename keywords.",
            "ATC-referenced certificates remain deferred until clause expansion.",
        ],
    }


def run(tender_id: str, tender_root: str | Path | None = None) -> int:
    output_paths = paths(tender_id, tender_root)
    output_paths["matrix_root"].mkdir(parents=True, exist_ok=True)
    manifest_rows = load_json(output_paths["manifest"])
    identity_rows = load_json(output_paths["identity"])
    requirements = load_json(output_paths["requirements"])

    matrix_rows = build_matrix_rows(manifest_rows, identity_rows, requirements)
    headers = CORE_HEADERS + requirement_column_groups(requirements) + REVIEWER_HEADERS
    issues = validate_rows(matrix_rows, requirements, identity_rows)
    notes = build_notes(tender_id, matrix_rows, requirements, headers, output_paths, issues)

    write_xlsx(output_paths["matrix_xlsx"], "Bidder Review Matrix", matrix_rows, headers)
    write_json(output_paths["matrix_json"], matrix_rows)
    write_json(output_paths["notes_json"], notes)
    write_xlsx(output_paths["notes_xlsx"], "Matrix Notes", [notes], NOTES_HEADERS)

    print(f"Matrix rows: {len(matrix_rows)}")
    print(f"Required document groups: {len(requirements)}")
    print(f"Matrix columns: {len(headers)}")
    print(f"Validation issues: {len(issues)}")
    for issue in issues:
        print(f"- {issue}")
    return 1 if issues else 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Build reviewer-facing bidder document matrix.")
    parser.add_argument("--tender-id", default=DEFAULT_TENDER_ID)
    parser.add_argument("--tender-root", default="", help="Optional explicit tender workspace root.")
    args = parser.parse_args()
    raise SystemExit(run(args.tender_id, args.tender_root or None))


if __name__ == "__main__":
    main()
