from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageEnhance, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pdf2image import convert_from_path
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TENDER_ID = "Tender_01_VacuumFlask"
TENDER_ID = DEFAULT_TENDER_ID
TENDER_ROOT = ROOT / DEFAULT_TENDER_ID
OUTPUT_ROOT = TENDER_ROOT / "05_Extraction_Output"
TEXT_ROOT = OUTPUT_ROOT / "01_text_extraction"
IDENTITY_ROOT = OUTPUT_ROOT / "02_identity_extraction"
VALIDATION_ROOT = OUTPUT_ROOT / "03_validation"
LOG_ROOT = OUTPUT_ROOT / "04_logs"

STATUS_EXTRACTED = "Extracted / Validated"
STATUS_NOT_FOUND = "Needs Review \u2014 not found by system"
STATUS_INVALID = "Needs Review \u2014 invalid candidate"
STATUS_MULTIPLE = "Needs Review \u2014 multiple candidates"
STATUS_OWNERSHIP = "Needs Review \u2014 ownership unclear"
STATUS_DERIVED = "Needs Review \u2014 derived from GSTIN"
STATUS_OCR_RETRY = "Needs Review \u2014 OCR retry used"
STATUS_FAILED = "Processing Failed"

PAN_RE = re.compile(r"(?<![A-Z0-9])([A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[A-Z])(?![A-Z0-9])")
GSTIN_RE = re.compile(r"(?<![A-Z0-9])([0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[A-Z](?:[\s\-.])*[0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[0-9](?:[\s\-.])*[A-Z](?:[\s\-.])*[1-9A-Z](?:[\s\-.])*Z(?:[\s\-.])*[0-9A-Z])(?![A-Z0-9])")
PAN_STRICT_RE = re.compile(r"^[A-Z]{5}[0-9]{4}[A-Z]$")
GSTIN_STRICT_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")
GSTIN_CHARSET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
VALID_GST_STATE_CODES = {
    "01", "02", "03", "04", "05", "06", "07", "08", "09", "10",
    "11", "12", "13", "14", "15", "16", "17", "18", "19", "20",
    "21", "22", "23", "24", "25", "26", "27", "29", "30", "31",
    "32", "33", "34", "35", "36", "37", "38", "97", "99",
}
_PADDLE_OCR: Any | None = None


def configure_tender_context(tender_id: str | None = None, tender_root: str | Path | None = None) -> None:
    global TENDER_ID, TENDER_ROOT, OUTPUT_ROOT, TEXT_ROOT, IDENTITY_ROOT, VALIDATION_ROOT, LOG_ROOT
    TENDER_ID = tender_id or DEFAULT_TENDER_ID
    TENDER_ROOT = Path(tender_root).resolve() if tender_root else ROOT / TENDER_ID
    OUTPUT_ROOT = TENDER_ROOT / "05_Extraction_Output"
    TEXT_ROOT = OUTPUT_ROOT / "01_text_extraction"
    IDENTITY_ROOT = OUTPUT_ROOT / "02_identity_extraction"
    VALIDATION_ROOT = OUTPUT_ROOT / "03_validation"
    LOG_ROOT = OUTPUT_ROOT / "04_logs"


@dataclass(frozen=True)
class DependencyStatus:
    paddleocr_available: bool
    tesseract_available: bool
    pdf_rendering_available: bool
    tesseract_path: str
    pdftoppm_path: str


def ensure_dirs() -> None:
    for folder in (
        TEXT_ROOT / "native_text",
        TEXT_ROOT / "ocr_text",
        TEXT_ROOT / "chosen_text",
        TEXT_ROOT / "page_json",
        IDENTITY_ROOT,
        VALIDATION_ROOT,
        LOG_ROOT,
    ):
        folder.mkdir(parents=True, exist_ok=True)


def log_event(event: dict[str, Any]) -> None:
    ensure_dirs()
    event = {"ts": time.strftime("%Y-%m-%dT%H:%M:%S%z"), **event}
    with (LOG_ROOT / "processing_log.jsonl").open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(event, ensure_ascii=False) + "\n")


def check_dependencies() -> DependencyStatus:
    try:
        import paddleocr  # noqa: F401

        paddleocr_available = True
    except Exception:
        paddleocr_available = False
    tesseract_path = shutil.which("tesseract") or ""
    pdftoppm_path = shutil.which("pdftoppm") or ""
    return DependencyStatus(
        paddleocr_available=paddleocr_available,
        tesseract_available=bool(tesseract_path),
        pdf_rendering_available=bool(pdftoppm_path),
        tesseract_path=tesseract_path,
        pdftoppm_path=pdftoppm_path,
    )


def bidder_from_path(path: Path) -> tuple[str, str]:
    parts = path.relative_to(TENDER_ROOT).parts
    for part in parts:
        if part.startswith("Bidder_"):
            bits = part.split("_", 2)
            bidder_id = "_".join(bits[:2])
            bidder_name = bits[2].replace("_", " ") if len(bits) > 2 else part
            return bidder_id, bidder_name
    return "", ""


def document_side(path: Path) -> str:
    rel = path.relative_to(TENDER_ROOT).as_posix()
    return "bidder-side" if rel.startswith("03_Bidder_Submissions/") else "tender-side"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def document_id(index: int, path: Path) -> str:
    short_hash = sha256_file(path)[:10].upper()
    return f"DOC-{index:04d}-{short_hash}"


def document_type_guess(path: Path) -> str:
    name = path.name.lower()
    if "gst" in name and "registration" in name:
        return "GST registration certificate"
    if "pan" in name:
        return "PAN card / PAN proof"
    if "udyam" in name or "msme" in name:
        return "Udyam / MSME certificate"
    if "mii" in name or "localcontent" in name or "local_content" in name:
        return "MII local content declaration"
    if "oem" in name and "turnover" in name:
        return "OEM turnover CA certificate"
    if "oem" in name and ("authorization" in name or "authorisation" in name):
        return "OEM authorization certificate"
    if "turnover" in name and "ca" in name:
        return "Bidder turnover CA certificate"
    if "integrity" in name or "declaration" in name or "annexure" in name:
        return "Bidder declaration / Integrity Pact"
    if "experience" in name or "crac" in name or "work" in name:
        return "Experience / CRAC / work order document"
    if "eligibility" in name or "atc" in name:
        return "ATC / eligibility criteria"
    if "biddoc" in name or "bid" in name:
        return "Tender bid document"
    return "Unclassified PDF"


def pdf_files() -> list[Path]:
    return sorted(
        path
        for path in TENDER_ROOT.rglob("*.pdf")
        if "05_Extraction_Output" not in path.parts
    )


def read_page_native_texts(path: Path) -> tuple[list[str], str]:
    try:
        reader = PdfReader(str(path))
        page_texts = []
        for page in reader.pages:
            try:
                page_texts.append(page.extract_text() or "")
            except Exception:
                page_texts.append("")
        return page_texts, ""
    except Exception as exc:
        return [], str(exc)


def run_tesseract(image_path: Path, tesseract_path: str, psm: str = "6") -> tuple[str, float | None, str]:
    cmd = [tesseract_path, str(image_path), "stdout", "-l", "eng", "--psm", psm, "tsv"]
    completed = subprocess.run(
        cmd,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=180,
    )
    if completed.returncode != 0:
        return "", None, completed.stderr.strip()
    lines = completed.stdout.splitlines()
    if not lines:
        return "", None, ""
    header = lines[0].split("\t")
    try:
        text_idx = header.index("text")
        conf_idx = header.index("conf")
        line_idx = header.index("line_num")
    except ValueError:
        return completed.stdout, None, "Unexpected Tesseract TSV header."
    words_by_line: dict[str, list[str]] = {}
    confs: list[float] = []
    for line in lines[1:]:
        cols = line.split("\t")
        if len(cols) <= max(text_idx, conf_idx, line_idx):
            continue
        word = cols[text_idx].strip()
        if not word:
            continue
        words_by_line.setdefault(cols[line_idx], []).append(word)
        try:
            conf = float(cols[conf_idx])
            if conf >= 0:
                confs.append(conf)
        except ValueError:
            pass
    text = "\n".join(" ".join(words) for _line, words in sorted(words_by_line.items()))
    avg_conf = round(sum(confs) / len(confs), 2) if confs else None
    return text, avg_conf, ""


def ocr_hit_score(text: str, conf: float | None) -> tuple[int, float, int]:
    pan_hits = len(PAN_RE.findall(text.upper()))
    gstin_hits = len(GSTIN_RE.findall(text.upper()))
    return (pan_hits * 20 + gstin_hits * 10, conf or 0.0, len(text.strip()))


def nonwhite_crop(image_path: Path) -> Image.Image:
    image = Image.open(image_path).convert("RGB")
    gray = ImageOps.grayscale(image)
    inverted = ImageOps.invert(gray)
    mask = inverted.point(lambda pixel: 255 if pixel > 25 else 0)
    bbox = mask.getbbox()
    if not bbox:
        return image
    x0, y0, x1, y1 = bbox
    pad = 40
    return image.crop(
        (
            max(0, x0 - pad),
            max(0, y0 - pad),
            min(image.width, x1 + pad),
            min(image.height, y1 + pad),
        )
    )


def preprocessed_variants(image_path: Path, output_dir: Path) -> list[Path]:
    crop = nonwhite_crop(image_path)
    variants: list[tuple[str, Image.Image]] = [("crop", crop)]
    gray = ImageOps.grayscale(crop)
    variants.append(("crop_gray", gray))
    high_contrast = ImageEnhance.Contrast(gray).enhance(2.2)
    variants.append(("crop_contrast_2x", high_contrast.resize((crop.width * 2, crop.height * 2))))
    paths: list[Path] = []
    for name, image in variants:
        path = output_dir / f"{name}.png"
        image.save(path)
        paths.append(path)
    return paths


def retry_tesseract_with_preprocessing(image_path: Path, tesseract_path: str, output_dir: Path) -> tuple[str, float | None, str]:
    best: tuple[str, float | None, str, tuple[int, float, int]] = ("", None, "", (0, 0.0, 0))
    for variant_path in preprocessed_variants(image_path, output_dir):
        for psm in ("6", "11", "12"):
            text, conf, err = run_tesseract(variant_path, tesseract_path, psm=psm)
            score = ocr_hit_score(text, conf)
            if score > best[3]:
                best = (text, conf, f"OCR retry used: {variant_path.stem}, psm {psm}. {err}".strip(), score)
    return best[0], best[1], best[2]


def _iter_paddle_pairs(value: Any) -> Iterable[tuple[str, float]]:
    if isinstance(value, tuple) and len(value) == 2 and isinstance(value[0], str):
        try:
            yield value[0], float(value[1])
        except Exception:
            yield value[0], 0.0
        return
    if isinstance(value, list):
        for item in value:
            yield from _iter_paddle_pairs(item)


def run_paddleocr(image_path: Path) -> tuple[str, float | None, str]:
    global _PADDLE_OCR
    try:
        from paddleocr import PaddleOCR

        if _PADDLE_OCR is None:
            _PADDLE_OCR = PaddleOCR(use_angle_cls=True, lang="en", show_log=False)
        result = _PADDLE_OCR.ocr(str(image_path), cls=True)
        pairs = list(_iter_paddle_pairs(result))
        text = "\n".join(pair[0] for pair in pairs if pair[0])
        confs = [pair[1] * 100 if pair[1] <= 1 else pair[1] for pair in pairs if pair[1] > 0]
        avg_conf = round(sum(confs) / len(confs), 2) if confs else None
        return text, avg_conf, ""
    except Exception as exc:
        return "", None, str(exc)


def ocr_page(path: Path, page_number: int, deps: DependencyStatus, engine_preference: str) -> tuple[str, str, float | None, str]:
    if not deps.pdf_rendering_available:
        return "", "failed", None, "PDF rendering unavailable."

    with tempfile.TemporaryDirectory() as tmp:
        images = convert_from_path(
            str(path),
            dpi=200,
            first_page=page_number,
            last_page=page_number,
            fmt="png",
            output_folder=tmp,
            paths_only=True,
        )
        if not images:
            return "", "failed", None, "PDF page rendering returned no image."
        image_path = Path(images[0])
        if engine_preference == "paddle" and deps.paddleocr_available:
            text, conf, err = run_paddleocr(image_path)
            if text.strip():
                return text, "paddleocr", conf, err
            fallback_note = f"PaddleOCR failed or returned no text; fallback used. Paddle note: {err}"
        else:
            fallback_note = "PaddleOCR unavailable; fallback used." if engine_preference == "paddle" else ""
        if deps.tesseract_available:
            text, conf, err = run_tesseract(image_path, deps.tesseract_path)
            retry_needed = not text.strip() or ocr_hit_score(text, conf)[0] == 0
            retry_note = ""
            if retry_needed:
                retry_text, retry_conf, retry_err = retry_tesseract_with_preprocessing(
                    image_path,
                    deps.tesseract_path,
                    Path(tmp),
                )
                if ocr_hit_score(retry_text, retry_conf) > ocr_hit_score(text, conf):
                    text, conf = retry_text, retry_conf
                    retry_note = retry_err or STATUS_OCR_RETRY
            notes = "; ".join(note for note in (fallback_note, err, retry_note) if note)
            return text, "tesseract", conf, notes
        return "", "failed", None, fallback_note or "No usable OCR engine found."


def choose_text(native_text: str, ocr_text: str) -> tuple[str, str]:
    native = native_text.strip()
    ocr = ocr_text.strip()
    if len(native) >= 80:
        if ocr and len(ocr) > len(native) * 1.5:
            return f"{native}\n\n{ocr}", "combined"
        return native, "native"
    if ocr:
        return ocr, "ocr"
    return native, "native"


def file_stem(document_id_value: str, file_name: str) -> str:
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(file_name).stem)[:80]
    return f"{document_id_value}_{safe_name}"


def write_text_outputs(stem: str, page_rows: list[dict[str, Any]]) -> None:
    native = "\n\n".join(f"--- Page {row['page_number']} ---\n{row['native_text']}" for row in page_rows)
    ocr = "\n\n".join(f"--- Page {row['page_number']} ---\n{row['ocr_text']}" for row in page_rows)
    chosen = "\n\n".join(f"--- Page {row['page_number']} ---\n{row['chosen_text']}" for row in page_rows)
    (TEXT_ROOT / "native_text" / f"{stem}.txt").write_text(native, encoding="utf-8")
    (TEXT_ROOT / "ocr_text" / f"{stem}.txt").write_text(ocr, encoding="utf-8")
    (TEXT_ROOT / "chosen_text" / f"{stem}.txt").write_text(chosen, encoding="utf-8")
    (TEXT_ROOT / "page_json" / f"{stem}.json").write_text(
        json.dumps(page_rows, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def build_manifest() -> tuple[list[dict[str, Any]], dict[str, list[str]]]:
    rows: list[dict[str, Any]] = []
    native_texts_by_doc: dict[str, list[str]] = {}
    for idx, path in enumerate(pdf_files(), start=1):
        doc_id = document_id(idx, path)
        bidder_id, bidder_name = bidder_from_path(path)
        native_texts, err = read_page_native_texts(path)
        page_count = len(native_texts)
        rows.append(
            {
                "tender_id": TENDER_ID,
                "bidder_id": bidder_id,
                "bidder_name": bidder_name,
                "document_id": doc_id,
                "document_side": document_side(path),
                "document_type_guess": document_type_guess(path),
                "file_name": path.name,
                "folder_path": path.parent.relative_to(TENDER_ROOT).as_posix(),
                "file_sha256": sha256_file(path),
                "file_size_bytes": path.stat().st_size,
                "page_count": page_count,
                "processing_status": "Pending" if not err else STATUS_FAILED,
                "error_message": err,
            }
        )
        native_texts_by_doc[doc_id] = native_texts
    return rows, native_texts_by_doc


def extract_pages(
    manifest_rows: list[dict[str, Any]],
    native_texts_by_doc: dict[str, list[str]],
    deps: DependencyStatus,
    engine_preference: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    page_rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for doc in manifest_rows:
        path = TENDER_ROOT / doc["folder_path"] / doc["file_name"]
        doc_page_rows: list[dict[str, Any]] = []
        if doc["processing_status"] == STATUS_FAILED:
            errors.append({"document_id": doc["document_id"], "file_name": doc["file_name"], "error_message": doc["error_message"]})
            continue
        for page_number, native_text in enumerate(native_texts_by_doc[doc["document_id"]], start=1):
            is_scanned_likely = len(native_text.strip()) < 40
            ocr_text = ""
            ocr_engine = "none"
            ocr_confidence_avg: float | None = None
            notes = ""
            if is_scanned_likely:
                try:
                    ocr_text, ocr_engine, ocr_confidence_avg, notes = ocr_page(path, page_number, deps, engine_preference)
                except Exception as exc:
                    ocr_engine = "failed"
                    notes = str(exc)
            chosen_text, chosen_source = choose_text(native_text, ocr_text)
            status = STATUS_EXTRACTED if chosen_text.strip() else STATUS_FAILED
            if status == STATUS_FAILED and not notes:
                notes = "No native or OCR text extracted."
            row = {
                "tender_id": TENDER_ID,
                "bidder_id": doc["bidder_id"],
                "document_id": doc["document_id"],
                "file_name": doc["file_name"],
                "page_number": page_number,
                "native_text": native_text,
                "native_text_chars": len(native_text.strip()),
                "ocr_text": ocr_text,
                "ocr_text_chars": len(ocr_text.strip()),
                "chosen_text": chosen_text,
                "chosen_text_source": chosen_source,
                "is_scanned_likely": is_scanned_likely,
                "ocr_engine": ocr_engine,
                "ocr_confidence_avg": ocr_confidence_avg,
                "processing_status": status,
                "processing_notes": notes,
            }
            page_rows.append(row)
            doc_page_rows.append(row)
            if status == STATUS_FAILED:
                errors.append({"document_id": doc["document_id"], "file_name": doc["file_name"], "page_number": page_number, "error_message": notes})
        stem = file_stem(doc["document_id"], doc["file_name"])
        write_text_outputs(stem, doc_page_rows)
        if doc_page_rows and all(row["processing_status"] == STATUS_EXTRACTED for row in doc_page_rows):
            doc["processing_status"] = "Processed"
            doc["error_message"] = ""
        elif doc_page_rows:
            doc["processing_status"] = STATUS_FAILED
            doc["error_message"] = "One or more pages failed text extraction."
    return page_rows, errors


def normalize_id(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def snippet(text: str, start: int, end: int, width: int = 80) -> str:
    left = max(start - width, 0)
    right = min(end + width, len(text))
    return re.sub(r"\s+", " ", text[left:right]).strip()


def gstin_checksum_valid(value: str) -> bool:
    if len(value) != 15:
        return False
    factor = 2
    total = 0
    for char in reversed(value[:14]):
        code_point = GSTIN_CHARSET.find(char)
        if code_point < 0:
            return False
        addend = factor * code_point
        factor = 1 if factor == 2 else 2
        addend = (addend // 36) + (addend % 36)
        total += addend
    check_code_point = (36 - (total % 36)) % 36
    return GSTIN_CHARSET[check_code_point] == value[14]


def gst_state_valid(value: str) -> bool:
    return len(value) >= 2 and value[:2] in VALID_GST_STATE_CODES


def owner_likelihood(document_side_value: str, doc_type: str, document_name: str) -> str:
    if document_side_value == "tender-side":
        return "unknown"
    haystack = f"{doc_type} {document_name}".lower()
    if any(token in haystack for token in ["gst registration", "pan card", "udyam", "msme", "firm registration", "declaration", "integrity"]):
        return "bidder_likely"
    if "ca certificate" in haystack or ("turnover" in haystack and "ca" in haystack):
        return "ca_firm_likely"
    if "oem" in haystack:
        return "oem_likely"
    if any(token in haystack for token in ["experience", "crac", "work order", "invoice"]):
        return "issuer_likely"
    return "unknown"


def candidate_confidence(regex_valid: bool, checksum_valid: bool | None, owner: str, source_method: str) -> float:
    score = 0.45
    if regex_valid:
        score += 0.2
    if checksum_valid is True:
        score += 0.15
    if owner == "bidder_likely":
        score += 0.15
    if source_method == "native":
        score += 0.05
    if source_method == "derived_from_gstin":
        score += 0.1
    return round(min(score, 0.99), 2)


def validation_status_for_candidate(field: str, regex_valid: bool, checksum_valid: bool | None, owner: str) -> str:
    if not regex_valid or (field == "GSTIN" and checksum_valid is not True):
        return STATUS_INVALID
    if owner != "bidder_likely":
        return STATUS_OWNERSHIP
    return STATUS_EXTRACTED


def extract_candidates(page_rows: list[dict[str, Any]], manifest_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    manifest_by_doc = {row["document_id"]: row for row in manifest_rows}
    candidates: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, int, str]] = set()
    for page in page_rows:
        doc = manifest_by_doc[page["document_id"]]
        text = page["chosen_text"] or ""
        for field_name, regex in (("PAN", PAN_RE), ("GSTIN", GSTIN_RE)):
            for match in regex.finditer(text.upper()):
                raw = match.group(1)
                normalized = normalize_id(raw)
                key = (page["document_id"], field_name, normalized, page["page_number"], page["chosen_text_source"])
                if key in seen:
                    continue
                seen.add(key)
                if field_name == "PAN":
                    regex_valid = bool(PAN_STRICT_RE.match(normalized))
                    checksum_valid = None
                    gstin_pan_match = None
                else:
                    regex_valid = bool(GSTIN_STRICT_RE.match(normalized)) and gst_state_valid(normalized)
                    checksum_valid = gstin_checksum_valid(normalized) if regex_valid else False
                    gstin_pan_match = "pan_unavailable"
                owner = owner_likelihood(doc["document_side"], doc["document_type_guess"], doc["file_name"])
                status = validation_status_for_candidate(field_name, regex_valid, checksum_valid, owner)
                candidates.append(
                    {
                        "tender_id": TENDER_ID,
                        "bidder_id": doc["bidder_id"],
                        "bidder_name": doc["bidder_name"],
                        "document_id": page["document_id"],
                        "document_name": doc["file_name"],
                        "document_type_guess": doc["document_type_guess"],
                        "document_side": doc["document_side"],
                        "field_name": field_name,
                        "extracted_value": raw,
                        "normalized_value": normalized,
                        "page": page["page_number"],
                        "evidence_text_snippet": snippet(text, match.start(1), match.end(1)),
                        "source_method": page["chosen_text_source"],
                        "ocr_engine": page["ocr_engine"],
                        "regex_valid": regex_valid,
                        "checksum_valid": checksum_valid,
                        "gstin_pan_match": gstin_pan_match,
                        "owner_likelihood": owner,
                        "confidence": candidate_confidence(regex_valid, checksum_valid, owner, page["chosen_text_source"]),
                        "validation_status": status,
                        "review_status": "Pending Review" if status != STATUS_EXTRACTED else "Not Required",
                        "notes": "",
                        "bbox_coordinates": None,
                    }
                )
    update_gstin_pan_matches(candidates)
    add_derived_pan_candidates(candidates)
    update_gstin_pan_matches(candidates)
    return candidates


def add_derived_pan_candidates(candidates: list[dict[str, Any]]) -> None:
    existing_keys = {
        (cand["bidder_id"], cand["normalized_value"], cand["document_id"], cand["page"], cand["source_method"])
        for cand in candidates
        if cand["field_name"] == "PAN"
    }
    derived: list[dict[str, Any]] = []
    for cand in candidates:
        if cand["field_name"] != "GSTIN":
            continue
        if cand["document_side"] != "bidder-side":
            continue
        if cand["owner_likelihood"] != "bidder_likely":
            continue
        if not cand["regex_valid"] or cand["checksum_valid"] is not True:
            continue
        pan_value = cand["normalized_value"][2:12]
        key = (cand["bidder_id"], pan_value, cand["document_id"], cand["page"], "derived_from_gstin")
        if key in existing_keys:
            continue
        existing_keys.add(key)
        derived.append(
            {
                "tender_id": cand["tender_id"],
                "bidder_id": cand["bidder_id"],
                "bidder_name": cand["bidder_name"],
                "document_id": cand["document_id"],
                "document_name": cand["document_name"],
                "document_type_guess": cand["document_type_guess"],
                "document_side": cand["document_side"],
                "field_name": "PAN",
                "extracted_value": pan_value,
                "normalized_value": pan_value,
                "page": cand["page"],
                "evidence_text_snippet": f"Derived from validated GSTIN {cand['normalized_value']}. Source snippet: {cand['evidence_text_snippet']}",
                "source_method": "derived_from_gstin",
                "ocr_engine": cand["ocr_engine"],
                "regex_valid": bool(PAN_STRICT_RE.match(pan_value)),
                "checksum_valid": None,
                "gstin_pan_match": None,
                "owner_likelihood": "bidder_likely",
                "confidence": candidate_confidence(True, None, "bidder_likely", "derived_from_gstin"),
                "validation_status": STATUS_DERIVED,
                "review_status": "Pending Review",
                "notes": "PAN derived from GSTIN characters 3-12; requires human/API confirmation.",
                "bbox_coordinates": None,
            }
        )
    candidates.extend(derived)


def update_gstin_pan_matches(candidates: list[dict[str, Any]]) -> None:
    pans_by_bidder: dict[str, set[str]] = {}
    for cand in candidates:
        if cand["document_side"] == "bidder-side" and cand["field_name"] == "PAN" and cand["regex_valid"]:
            pans_by_bidder.setdefault(cand["bidder_id"], set()).add(cand["normalized_value"])
    for cand in candidates:
        if cand["field_name"] != "GSTIN" or cand["document_side"] != "bidder-side":
            continue
        bidder_pans = pans_by_bidder.get(cand["bidder_id"], set())
        if not bidder_pans:
            cand["gstin_pan_match"] = "pan_unavailable"
        elif cand["normalized_value"][2:12] in bidder_pans:
            cand["gstin_pan_match"] = "match"
        else:
            cand["gstin_pan_match"] = "mismatch"


def high_trust(cand: dict[str, Any]) -> bool:
    return cand["document_side"] == "bidder-side" and cand["owner_likelihood"] == "bidder_likely"


def preferred_pan(pan_candidates: list[dict[str, Any]], gstin_candidates: list[dict[str, Any]]) -> tuple[dict[str, Any] | None, str]:
    direct_candidates = [cand for cand in pan_candidates if cand.get("source_method") != "derived_from_gstin"]
    derived_candidates = [cand for cand in pan_candidates if cand.get("source_method") == "derived_from_gstin"]
    valid_high = [cand for cand in direct_candidates if cand["regex_valid"] and high_trust(cand)]
    if not valid_high:
        valid_derived = [cand for cand in derived_candidates if cand["regex_valid"] and high_trust(cand)]
        unique_derived = sorted({cand["normalized_value"] for cand in valid_derived})
        if len(unique_derived) == 1:
            return valid_derived[0], STATUS_DERIVED
        if len(unique_derived) > 1:
            return None, STATUS_MULTIPLE
        valid_low = [cand for cand in direct_candidates if cand["regex_valid"] and cand["document_side"] == "bidder-side"]
        if valid_low:
            return None, STATUS_OWNERSHIP
        return None, STATUS_NOT_FOUND
    gstin_embedded_pans = {cand["normalized_value"][2:12] for cand in gstin_candidates if cand["regex_valid"] and cand["checksum_valid"] is True}
    supported = [cand for cand in valid_high if cand["normalized_value"] in gstin_embedded_pans]
    pool = supported or valid_high
    unique_values = sorted({cand["normalized_value"] for cand in pool})
    if len(unique_values) > 1:
        return None, STATUS_MULTIPLE
    return pool[0], STATUS_EXTRACTED


def preferred_gstin(gstin_candidates: list[dict[str, Any]], preferred_pan_value: str | None) -> tuple[dict[str, Any] | None, str]:
    valid_high = [
        cand
        for cand in gstin_candidates
        if cand["regex_valid"] and cand["checksum_valid"] is True and high_trust(cand)
    ]
    if preferred_pan_value:
        matching = [cand for cand in valid_high if cand["normalized_value"][2:12] == preferred_pan_value]
        if matching:
            unique_values = sorted({cand["normalized_value"] for cand in matching})
            if len(unique_values) > 1:
                return None, STATUS_MULTIPLE
            return matching[0], STATUS_EXTRACTED
    if not valid_high:
        valid_low = [
            cand
            for cand in gstin_candidates
            if cand["regex_valid"] and cand["checksum_valid"] is True and cand["document_side"] == "bidder-side"
        ]
        if valid_low:
            return None, STATUS_OWNERSHIP
        invalid = [cand for cand in gstin_candidates if cand["document_side"] == "bidder-side"]
        if invalid:
            return None, STATUS_INVALID
        return None, STATUS_NOT_FOUND
    unique_values = sorted({cand["normalized_value"] for cand in valid_high})
    if len(unique_values) > 1:
        return None, STATUS_MULTIPLE
    return valid_high[0], STATUS_EXTRACTED


def build_identity_summary(manifest_rows: list[dict[str, Any]], candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    bidders = sorted({(row["bidder_id"], row["bidder_name"]) for row in manifest_rows if row["bidder_id"]})
    rows: list[dict[str, Any]] = []
    for bidder_id, bidder_name in bidders:
        bidder_candidates = [cand for cand in candidates if cand["bidder_id"] == bidder_id and cand["document_side"] == "bidder-side"]
        pan_candidates = [cand for cand in bidder_candidates if cand["field_name"] == "PAN"]
        gstin_candidates = [cand for cand in bidder_candidates if cand["field_name"] == "GSTIN"]
        pan, pan_status = preferred_pan(pan_candidates, gstin_candidates)
        gstin, gstin_status = preferred_gstin(gstin_candidates, pan["normalized_value"] if pan else None)
        if pan and gstin:
            match_status = "match" if gstin["normalized_value"][2:12] == pan["normalized_value"] else "mismatch"
        else:
            match_status = "not_tested"
        ambiguity = "Yes" if pan_status == STATUS_MULTIPLE or gstin_status == STATUS_MULTIPLE else "No"
        needs_review = "No" if pan_status == STATUS_EXTRACTED and gstin_status == STATUS_EXTRACTED and match_status == "match" else "Yes"
        rows.append(
            {
                "tender_id": TENDER_ID,
                "bidder_id": bidder_id,
                "bidder_name": bidder_name,
                "pan": pan["normalized_value"] if pan else "",
                "pan_source_file": pan["document_name"] if pan else "",
                "pan_page": pan["page"] if pan else "",
                "pan_status": pan_status,
                "gstin": gstin["normalized_value"] if gstin else "",
                "gstin_source_file": gstin["document_name"] if gstin else "",
                "gstin_page": gstin["page"] if gstin else "",
                "gstin_status": gstin_status,
                "gstin_pan_match_status": match_status,
                "pan_candidate_count": len(pan_candidates),
                "gstin_candidate_count": len(gstin_candidates),
                "ambiguity_flag": ambiguity,
                "needs_human_review": needs_review,
                "reviewer_decision": "",
                "reviewer_notes": "",
                "pan_api_status_reserved": "",
                "gstin_api_status_reserved": "",
                "api_verified_at_reserved": "",
                "api_source_reserved": "",
            }
        )
    return rows


def validation_rows(candidates: list[dict[str, Any]], identity_summary: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for cand in candidates:
        rows.append(
            {
                "record_type": "candidate",
                "tender_id": cand["tender_id"],
                "bidder_id": cand["bidder_id"],
                "field_name": cand["field_name"],
                "normalized_value": cand["normalized_value"],
                "regex_valid": cand["regex_valid"],
                "checksum_valid": cand["checksum_valid"],
                "gstin_pan_match": cand["gstin_pan_match"],
                "owner_likelihood": cand["owner_likelihood"],
                "validation_status": cand["validation_status"],
                "source_file": cand["document_name"],
                "page": cand["page"],
                "notes": cand["notes"],
            }
        )
    for row in identity_summary:
        rows.append(
            {
                "record_type": "bidder_summary",
                "tender_id": row["tender_id"],
                "bidder_id": row["bidder_id"],
                "field_name": "PAN/GSTIN",
                "normalized_value": (
                    f"PAN={row.get('pan', row.get('preferred_pan', ''))}; "
                    f"GSTIN={row.get('gstin', row.get('preferred_gstin', ''))}"
                ),
                "regex_valid": "",
                "checksum_valid": "",
                "gstin_pan_match": row["gstin_pan_match_status"],
                "owner_likelihood": "",
                "validation_status": STATUS_EXTRACTED if row["needs_human_review"] == "No" else "Needs Review",
                "source_file": "",
                "page": "",
                "notes": f"Needs human review: {row['needs_human_review']}",
            }
        )
    return rows


def autosize(ws) -> None:
    for column_cells in ws.columns:
        col = get_column_letter(column_cells[0].column)
        length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            length = max(length, len(value))
        ws.column_dimensions[col].width = min(max(length + 2, 12), 60)


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)
    wb.save(path)


def write_json(path: Path, rows: list[dict[str, Any]]) -> None:
    path.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")


def write_outputs(
    manifest_rows: list[dict[str, Any]],
    page_rows: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
    summary: list[dict[str, Any]],
    validations: list[dict[str, Any]],
    errors: list[dict[str, Any]],
) -> None:
    manifest_headers = [
        "tender_id", "bidder_id", "bidder_name", "document_id", "document_side",
        "document_type_guess", "file_name", "folder_path", "file_sha256", "file_size_bytes",
        "page_count", "processing_status", "error_message",
    ]
    candidate_headers = [
        "tender_id", "bidder_id", "bidder_name", "document_id", "document_name",
        "document_type_guess", "document_side", "field_name", "extracted_value",
        "normalized_value", "page", "evidence_text_snippet", "source_method", "ocr_engine",
        "regex_valid", "checksum_valid", "gstin_pan_match", "owner_likelihood", "confidence",
        "validation_status", "review_status", "notes", "bbox_coordinates",
    ]
    summary_headers = [
        "tender_id", "bidder_id", "bidder_name", "pan", "pan_source_file",
        "pan_page", "pan_status", "gstin", "gstin_source_file",
        "gstin_page", "gstin_status", "gstin_pan_match_status", "pan_candidate_count",
        "gstin_candidate_count", "ambiguity_flag", "needs_human_review", "reviewer_decision",
        "reviewer_notes", "pan_api_status_reserved", "gstin_api_status_reserved",
        "api_verified_at_reserved", "api_source_reserved",
    ]
    validation_headers = [
        "record_type", "tender_id", "bidder_id", "field_name", "normalized_value",
        "regex_valid", "checksum_valid", "gstin_pan_match", "owner_likelihood",
        "validation_status", "source_file", "page", "notes",
    ]
    error_headers = ["document_id", "file_name", "page_number", "error_message"]
    write_xlsx(OUTPUT_ROOT / "document_manifest.xlsx", "Document Manifest", manifest_rows, manifest_headers)
    write_json(OUTPUT_ROOT / "document_manifest.json", manifest_rows)
    write_json(TEXT_ROOT / "page_json" / "all_pages.json", page_rows)
    write_xlsx(IDENTITY_ROOT / "pan_gstin_candidates.xlsx", "PAN GSTIN Candidates", candidates, candidate_headers)
    write_json(IDENTITY_ROOT / "pan_gstin_candidates.json", candidates)
    write_xlsx(IDENTITY_ROOT / "bidder_identity_summary.xlsx", "Bidder Identity Summary", summary, summary_headers)
    write_json(IDENTITY_ROOT / "bidder_identity_summary.json", summary)
    write_xlsx(VALIDATION_ROOT / "validation_results.xlsx", "Validation Results", validations, validation_headers)
    write_json(VALIDATION_ROOT / "validation_results.json", validations)
    write_xlsx(LOG_ROOT / "errors.xlsx", "Errors", errors, error_headers)


def dependency_report(deps: DependencyStatus) -> list[dict[str, str]]:
    return [
        {"dependency": "PaddleOCR available", "status": "Yes" if deps.paddleocr_available else "No", "detail": ""},
        {"dependency": "Tesseract available", "status": "Yes" if deps.tesseract_available else "No", "detail": deps.tesseract_path},
        {"dependency": "PDF rendering available", "status": "Yes" if deps.pdf_rendering_available else "No", "detail": deps.pdftoppm_path},
    ]


def validate_acceptance(
    manifest_rows: list[dict[str, Any]],
    page_rows: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
    summary: list[dict[str, Any]],
) -> list[str]:
    issues: list[str] = []
    if len(manifest_rows) != len(pdf_files()):
        issues.append("Manifest PDF count does not match source PDF count.")
    pages_by_doc = {row["document_id"] for row in page_rows}
    for doc in manifest_rows:
        if doc["processing_status"] != "Processed" and doc["document_id"] not in pages_by_doc:
            issues.append(f"No page rows for {doc['document_id']} {doc['file_name']}.")
    for cand in candidates:
        if not cand["document_name"] or not cand["page"]:
            issues.append(f"Candidate missing source/page: {cand['normalized_value']}.")
        if not cand["evidence_text_snippet"]:
            issues.append(f"Candidate missing evidence snippet: {cand['normalized_value']}.")
        if cand["document_side"] == "tender-side" and cand["owner_likelihood"] == "bidder_likely":
            issues.append(f"Tender-side candidate marked bidder_likely: {cand['normalized_value']}.")
    bidder_count = len({row["bidder_id"] for row in manifest_rows if row["bidder_id"]})
    if len(summary) != bidder_count:
        issues.append("Bidder identity summary row count does not match bidder count.")
    forbidden = ["rejected", "disqualified", "accepted", "non-compliant"]
    for row in summary:
        text = json.dumps(row, ensure_ascii=False).lower()
        if any(term in text for term in forbidden):
            issues.append(f"Forbidden decision wording in summary for {row['bidder_id']}.")
    return issues


def run(
    engine: str,
    check_only: bool = False,
    tender_id: str | None = None,
    tender_root: str | Path | None = None,
) -> int:
    configure_tender_context(tender_id, tender_root)
    ensure_dirs()
    deps = check_dependencies()
    dep_rows = dependency_report(deps)
    write_xlsx(LOG_ROOT / "dependency_check.xlsx", "Dependency Check", dep_rows, ["dependency", "status", "detail"])
    write_json(LOG_ROOT / "dependency_check.json", dep_rows)
    print("Dependency check:")
    for row in dep_rows:
        print(f"- {row['dependency']}: {row['status']} {row['detail']}")
    if check_only:
        return 0
    log_event({"event": "start", "engine_requested": engine})
    manifest_rows, native_texts_by_doc = build_manifest()
    page_rows, errors = extract_pages(manifest_rows, native_texts_by_doc, deps, engine)
    candidates = extract_candidates(page_rows, manifest_rows)
    summary = build_identity_summary(manifest_rows, candidates)
    validations = validation_rows(candidates, summary)
    issues = validate_acceptance(manifest_rows, page_rows, candidates, summary)
    for issue in issues:
        errors.append({"document_id": "", "file_name": "", "page_number": "", "error_message": f"ACCEPTANCE: {issue}"})
    write_outputs(manifest_rows, page_rows, candidates, summary, validations, errors)
    log_event(
        {
            "event": "complete",
            "documents": len(manifest_rows),
            "pages": len(page_rows),
            "candidates": len(candidates),
            "bidders": len(summary),
            "acceptance_issues": len(issues),
        }
    )
    print(f"Processed documents: {len(manifest_rows)}")
    print(f"Processed pages: {len(page_rows)}")
    print(f"PAN/GSTIN candidates: {len(candidates)}")
    print(f"Bidder summaries: {len(summary)}")
    print(f"Acceptance issues: {len(issues)}")
    return 1 if issues else 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Run Phase 1 OCR/text and PAN/GSTIN extraction.")
    parser.add_argument("--engine", choices=["auto", "paddle", "tesseract"], default="auto")
    parser.add_argument("--check-deps", action="store_true", help="Only write dependency report and exit.")
    parser.add_argument("--tender-id", default=DEFAULT_TENDER_ID)
    parser.add_argument("--tender-root", default="", help="Optional explicit tender workspace root.")
    args = parser.parse_args()
    engine = args.engine
    if engine == "auto":
        engine = "paddle"
    raise SystemExit(run(engine, args.check_deps, args.tender_id, args.tender_root or None))


if __name__ == "__main__":
    main()
