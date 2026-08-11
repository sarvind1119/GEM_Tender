"""Verification orchestration: targets -> providers -> outputs + audit trail.

Outputs under ``<tender_root>/05_Extraction_Output/08_verification/``:

- ``verification_audit.jsonl``   append-only, one line per registry call
- ``verification_matrix.json``   latest result per (bidder, service), atomic write
- ``verification_matrix.xlsx``   reviewer-facing export (same styling as other tools)

Additionally fills the reserved fields in ``bidder_identity_summary.json``
(``pan_api_status_reserved``, ``gstin_api_status_reserved``,
``api_verified_at_reserved``, ``api_source_reserved``) and upgrades
``needs_human_review`` to "Yes" when any outcome is not ``verified``
(never downgrades). Registry outcomes are reviewer evidence only.
"""

from __future__ import annotations

import json
import time
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.verification import cin_extraction
from tools.verification.providers import PARSERS, get_provider
from tools.verification.schema import (
    ALL_SERVICES,
    INPUT_PATTERNS,
    OUTCOME_SKIPPED,
    OUTCOME_VERIFIED,
    SERVICE_GSTIN,
    SERVICE_MCA,
    SERVICE_PAN,
    VerificationResult,
    review_note_for,
)

MATRIX_HEADERS = [
    "bidder_id",
    "bidder_name",
    "service",
    "input_value",
    "outcome",
    "registry_status",
    "registered_name",
    "trade_name",
    "name_match_score",
    "name_match_band",
    "review_note",
    "provider",
    "provider_ref",
    "verified_at",
    "raw_response_sha256",
    "extra",
]


def now_iso() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%S%z")


def verification_root(tender_root: Path) -> Path:
    return Path(tender_root) / "05_Extraction_Output" / "08_verification"


def identity_summary_path(tender_root: Path) -> Path:
    return (
        Path(tender_root)
        / "05_Extraction_Output"
        / "02_identity_extraction"
        / "bidder_identity_summary.json"
    )


def matrix_json_path(tender_root: Path) -> Path:
    return verification_root(tender_root) / "verification_matrix.json"


def matrix_xlsx_path(tender_root: Path) -> Path:
    return verification_root(tender_root) / "verification_matrix.xlsx"


def audit_log_path(tender_root: Path) -> Path:
    return verification_root(tender_root) / "verification_audit.jsonl"


def write_json_atomic(path: Path, payload: Any) -> None:
    # Local copy of the runner's helper to avoid import coupling with
    # tools/pipeline_runner.py (edited in parallel).
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def append_audit(path: Path, record: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
        handle.flush()


def load_identity_rows(tender_root: Path) -> list[dict[str, Any]]:
    path = identity_summary_path(tender_root)
    if not path.exists():
        return []
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []


def identity_value(row: dict[str, Any], field: str) -> str:
    # Identity summary rows have used both `pan` and `preferred_pan` naming.
    return str(row.get(field, row.get(f"preferred_{field}", "")) or "").strip().upper()


def _skip_result(
    service: str, bidder_id: str, bidder_name: str, input_value: str, detail: str
) -> VerificationResult:
    result = VerificationResult(
        service=service,
        input_value=input_value,
        bidder_id=bidder_id,
        bidder_name=bidder_name,
        outcome=OUTCOME_SKIPPED,
        provider="none",
        verified_at=now_iso(),
        extra={"detail": detail},
    )
    result.review_note = review_note_for(result)
    return result


def resolve_mca_input(
    bidder_id: str,
    cin_map: dict[str, str] | None,
    extracted: dict[str, list[dict[str, Any]]],
) -> tuple[str, str]:
    """Return (cin, source_detail). Manual entry wins; extraction must be unambiguous."""
    manual = ((cin_map or {}).get(bidder_id) or "").strip().upper()
    if manual:
        return manual, "reviewer_entry"
    candidates = extracted.get(bidder_id, [])
    if not candidates:
        return "", "no CIN found in bidder documents; enter it manually to enable the MCA check"
    unique = cin_extraction.unique_cin_for_bidder(candidates)
    if unique:
        return unique, "document_extraction"
    distinct = sorted({row["cin"] for row in candidates})
    return "", (
        f"multiple distinct CIN candidates found ({', '.join(distinct)}); "
        "reviewer selection needed"
    )


def verify_one(
    service: str,
    input_value: str,
    bidder_id: str,
    bidder_name: str,
    *,
    audit_path: Path,
    triggered_by: str,
) -> VerificationResult:
    provider = get_provider(service)
    call = {
        SERVICE_PAN: provider.verify_pan,
        SERVICE_GSTIN: provider.verify_gstin,
        SERVICE_MCA: provider.verify_cin,
    }[service]
    started = time.perf_counter()
    raw = call(input_value, bidder_name)
    duration_ms = round((time.perf_counter() - started) * 1000, 1)
    result = PARSERS[service](
        raw,
        provider.name,
        bidder_id=bidder_id,
        bidder_name=bidder_name,
        input_value=input_value,
    )
    append_audit(
        audit_path,
        {
            "ts": now_iso(),
            "service": service,
            "provider": provider.name,
            "bidder_id": bidder_id,
            "input_value": input_value,
            "http_status": raw[0],
            "raw_response": raw[1],
            "outcome": result.outcome,
            "provider_ref": result.provider_ref,
            "duration_ms": duration_ms,
            "triggered_by": triggered_by,
        },
    )
    return result


def run_verification(
    tender_root: str | Path,
    *,
    bidder_ids: list[str] | None = None,
    services: tuple[str, ...] | list[str] = ALL_SERVICES,
    cin_map: dict[str, str] | None = None,
    triggered_by: str = "cli",
    log: Callable[[str], Any] = print,
) -> dict[str, Any]:
    tender_root = Path(tender_root)
    identity_rows = load_identity_rows(tender_root)
    if not identity_rows:
        log("No bidder identity summary found; run the pipeline (Agent 1) first.")
        return {"results": 0, "error": "identity summary missing"}

    services = [service for service in services if service in ALL_SERVICES]
    audit_path = audit_log_path(tender_root)
    extracted_cins = (
        cin_extraction.cin_candidates_by_bidder(tender_root) if SERVICE_MCA in services else {}
    )
    results: list[VerificationResult] = []

    for row in identity_rows:
        bidder_id = str(row.get("bidder_id", ""))
        bidder_name = str(row.get("bidder_name", ""))
        if bidder_ids and bidder_id not in bidder_ids:
            continue
        inputs = {
            SERVICE_PAN: identity_value(row, "pan"),
            SERVICE_GSTIN: identity_value(row, "gstin"),
        }
        for service in services:
            if service == SERVICE_MCA:
                value, detail = resolve_mca_input(bidder_id, cin_map, extracted_cins)
                if not value:
                    results.append(_skip_result(service, bidder_id, bidder_name, "", detail))
                    continue
                source_detail = detail
            else:
                value = inputs[service]
                source_detail = "identity_summary"
                if not value:
                    results.append(
                        _skip_result(
                            service, bidder_id, bidder_name, "",
                            "no extracted value available in the identity summary",
                        )
                    )
                    continue
            if not INPUT_PATTERNS[service].match(value):
                results.append(
                    _skip_result(
                        service, bidder_id, bidder_name, value,
                        "value does not match the expected format; not sent to the registry",
                    )
                )
                continue
            log(f"Verifying {service.upper()} for {bidder_id}...")
            result = verify_one(
                service, value, bidder_id, bidder_name,
                audit_path=audit_path, triggered_by=triggered_by,
            )
            result.extra.setdefault("input_source", source_detail)
            results.append(result)

    write_outputs(tender_root, results)
    update_reserved_fields(tender_root, results)
    summary = summarize(results)
    log(
        f"Verification finished: {summary['verified']} verified, "
        f"{summary['flagged']} flagged, {summary['not_found']} not found, "
        f"{summary['errors']} errors, {summary['skipped']} skipped."
    )
    return summary


def summarize(results: list[VerificationResult]) -> dict[str, Any]:
    def count(outcome: str) -> int:
        return sum(1 for result in results if result.outcome == outcome)

    return {
        "results": len(results),
        "verified": count("verified"),
        "flagged": count("status_flag"),
        "not_found": count("not_found"),
        "errors": count("error"),
        "skipped": count("skipped"),
    }


def excel_safe(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", value)
    if isinstance(value, (dict, list)):
        return ILLEGAL_CHARACTERS_RE.sub(" ", json.dumps(value, ensure_ascii=False))
    return value


def write_matrix_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Registry Verification"
    ws.append(MATRIX_HEADERS)
    for row in rows:
        ws.append([excel_safe(row.get(header, "")) for header in MATRIX_HEADERS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for excel_row in ws.iter_rows(min_row=2):
        for cell in excel_row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column_cells in ws.columns:
        column = get_column_letter(column_cells[0].column)
        max_len = max(len("" if cell.value is None else str(cell.value)) for cell in column_cells)
        ws.column_dimensions[column].width = min(max(max_len + 2, 12), 60)
    wb.save(path)


def write_outputs(tender_root: Path, results: list[VerificationResult]) -> None:
    # Keep the latest result per (bidder, service): merge new results over any
    # previous matrix so partial runs (one bidder / one service) do not erase
    # earlier findings.
    previous: dict[tuple[str, str], dict[str, Any]] = {}
    path = matrix_json_path(tender_root)
    if path.exists():
        try:
            for row in json.loads(path.read_text(encoding="utf-8")):
                previous[(row.get("bidder_id", ""), row.get("service", ""))] = row
        except (json.JSONDecodeError, OSError):
            previous = {}
    for result in results:
        previous[(result.bidder_id, result.service)] = result.to_row()
    rows = sorted(previous.values(), key=lambda row: (row.get("bidder_id", ""), row.get("service", "")))
    write_json_atomic(path, rows)
    write_matrix_xlsx(matrix_xlsx_path(tender_root), rows)


def update_reserved_fields(tender_root: Path, results: list[VerificationResult]) -> None:
    path = identity_summary_path(tender_root)
    rows = load_identity_rows(tender_root)
    if not rows:
        return
    by_bidder: dict[str, dict[str, VerificationResult]] = {}
    for result in results:
        by_bidder.setdefault(result.bidder_id, {})[result.service] = result
    for row in rows:
        bidder_results = by_bidder.get(str(row.get("bidder_id", "")))
        if not bidder_results:
            continue
        pan_result = bidder_results.get(SERVICE_PAN)
        gstin_result = bidder_results.get(SERVICE_GSTIN)
        if pan_result and pan_result.outcome != OUTCOME_SKIPPED:
            row["pan_api_status_reserved"] = f"{pan_result.outcome}: {pan_result.registry_status}".strip(": ")
        if gstin_result and gstin_result.outcome != OUTCOME_SKIPPED:
            row["gstin_api_status_reserved"] = f"{gstin_result.outcome}: {gstin_result.registry_status}".strip(": ")
        acted = [r for r in bidder_results.values() if r.outcome != OUTCOME_SKIPPED]
        if acted:
            row["api_verified_at_reserved"] = max(r.verified_at for r in acted)
            row["api_source_reserved"] = "; ".join(
                sorted({f"{r.provider}:{r.service}" for r in acted})
            )
            # Evidence-only escalation: never downgrade an existing "Yes".
            if any(r.outcome != OUTCOME_VERIFIED for r in acted):
                row["needs_human_review"] = "Yes"
    write_json_atomic(path, rows)
