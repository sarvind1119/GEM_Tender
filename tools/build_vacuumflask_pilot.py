from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
TENDER_ID = "Tender_01_VacuumFlask"
TENDER_ROOT = ROOT / TENDER_ID
CONTROL_DIR = TENDER_ROOT / "04_Control_Artifacts"
EXTRACTION_DIR = TENDER_ROOT / "05_Extraction_Output"
MATRIX_DIR = TENDER_ROOT / "06_Compliance_Matrix"
AUDIT_DIR = TENDER_ROOT / "07_Audit_Log"

PAN_RE = re.compile(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b")
GSTIN_RE = re.compile(r"\b[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b")


@dataclass(frozen=True)
class Requirement:
    req_id: str
    clause_ref: str
    requirement_text: str
    required_document: str
    field_to_extract: str
    validation_rule: str
    tier: str
    decision_rule: str
    remarks: str


REQUIREMENTS = [
    Requirement(
        "T1-R01",
        "Pilot control",
        "Bidder identity / firm name must be traceable to a bidder submission folder.",
        "Bidder submission folder",
        "Bidder name",
        "Folder exists and contains submitted PDFs.",
        "1",
        "Mark Present only when bidder folder contains at least one PDF.",
        "Identity source is folder name for the pilot; human to confirm legal name.",
    ),
    Requirement(
        "T1-R02",
        "Tender/ATC identity checks",
        "PAN should be captured when present in submitted documents.",
        "PAN proof or PAN-bearing certificate",
        "PAN",
        "PAN regex AAAAA9999A.",
        "1",
        "Mark Present only on exact PAN-format extraction; otherwise Needs Review.",
        "No automatic rejection if PAN is not found before OCR.",
    ),
    Requirement(
        "T1-R03",
        "Tender/ATC identity checks",
        "GSTIN should be captured when present in submitted documents.",
        "GST registration certificate",
        "GSTIN",
        "GSTIN regex and offline GSTIN-PAN cross-check where PAN is available.",
        "1",
        "Mark Present only on exact GSTIN-format extraction; otherwise Needs Review.",
        "External GST API verification is out of scope until approval.",
    ),
    Requirement(
        "T1-R04",
        "Tender/ATC identity checks",
        "GST registration certificate must be present when submitted by bidder.",
        "GST registration certificate",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when GST registration PDF is identified.",
        "Presence only; validity remains human-reviewable until rules are approved.",
    ),
    Requirement(
        "T1-R05",
        "Tender/ATC identity checks",
        "GSTIN and PAN should cross-match when both are available.",
        "GST registration certificate and PAN proof",
        "GSTIN chars 3-12 vs PAN",
        "GSTIN positions 3-12 must equal PAN.",
        "1",
        "Mark Compliant only when both values exist and match exactly.",
        "If either value is missing, route to Needs Review.",
    ),
    Requirement(
        "T1-R06",
        "Bid document - Experience and Turnover",
        "Bidder turnover CA certificate must be present.",
        "Bidder turnover CA certificate",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when bidder turnover certificate PDF is identified.",
        "Presence only; turnover amount comparison is held back.",
    ),
    Requirement(
        "T1-R07",
        "Bid document - OEM authorization",
        "OEM authorization certificate must be present.",
        "OEM authorization certificate",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when OEM authorization PDF is identified.",
        "Presence only.",
    ),
    Requirement(
        "T1-R08",
        "Bid document - OEM annual turnover",
        "OEM turnover certificate must be present where applicable.",
        "OEM turnover CA certificate",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when OEM turnover certificate PDF is identified.",
        "If not identified, route to review because applicability may vary.",
    ),
    Requirement(
        "T1-R09",
        "MII purchase preference",
        "MII / local content declaration must be present where claimed or required.",
        "MII local content declaration",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when MII/local content PDF is identified.",
        "Presence only; content percentage review is held back.",
    ),
    Requirement(
        "T1-R10",
        "ATC declarations",
        "Integrity Pact / required declarations must be present.",
        "Integrity Pact or declaration document",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when integrity/declaration PDF is identified.",
        "No rejection on absence; route to human review.",
    ),
    Requirement(
        "T1-R11",
        "MSME / Udyam preference",
        "Udyam / MSME certificate must be present where applicable.",
        "Udyam / MSME certificate",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when Udyam/MSME PDF is identified; otherwise review applicability.",
        "Preference applicability to be confirmed manually.",
    ),
    Requirement(
        "T1-R12",
        "ATC certificates",
        "Type test / BIS / ISO certificates must be located if submitted.",
        "BIS / ISO / type test certificate",
        "Document presence",
        "Expected document type exists or keywords appear in extracted text.",
        "1",
        "Mark Present only when matching document or text evidence is found.",
        "Presence only; product compliance is Tier 2/future review.",
    ),
    Requirement(
        "T1-R13",
        "ATC experience criteria",
        "Experience / CRAC / work order documents must be present.",
        "Experience / CRAC / work order document",
        "Document presence",
        "Expected document type exists in bidder folder.",
        "1",
        "Mark Present when experience/work-order PDF is identified.",
        "Only presence is checked; eligibility match is held back.",
    ),
]


def ensure_dirs() -> None:
    for directory in (CONTROL_DIR, EXTRACTION_DIR, MATRIX_DIR, AUDIT_DIR):
        directory.mkdir(parents=True, exist_ok=True)


def bidder_from_path(path: Path) -> tuple[str, str]:
    parts = path.relative_to(TENDER_ROOT).parts
    for part in parts:
        if part.startswith("Bidder_"):
            bits = part.split("_", 2)
            bidder_id = "_".join(bits[:2])
            bidder_name = bits[2].replace("_", " ") if len(bits) > 2 else part
            return bidder_id, bidder_name
    return "", ""


def document_category(path: Path) -> str:
    rel = path.relative_to(TENDER_ROOT).as_posix()
    if rel.startswith("01_Tender_Document/"):
        return "Tender Document"
    if rel.startswith("02_Eligibility_and_Terms/"):
        return "Eligibility and Terms"
    if rel.startswith("03_Bidder_Submissions/"):
        return "Bidder Submission"
    return "Other"


def expected_doc_type(path: Path, text_hint: str = "") -> str:
    name = path.name.lower()
    haystack = f"{name} {text_hint[:5000].lower()}"
    if "gst" in haystack and "registration" in haystack:
        return "GST registration certificate"
    if "mii" in haystack or "localcontent" in haystack or "local content" in haystack:
        return "MII local content declaration"
    if "oem" in haystack and "turnover" in haystack:
        return "OEM turnover CA certificate"
    if "oem" in haystack and ("authorization" in haystack or "authorisation" in haystack):
        return "OEM authorization certificate"
    if "turnover" in haystack and "ca" in haystack:
        return "Bidder turnover CA certificate"
    if "integrity" in haystack or "declaration" in haystack or "annexure" in haystack:
        return "Integrity Pact or declaration document"
    if "udyam" in haystack or "msme" in haystack:
        return "Udyam / MSME certificate"
    if "experience" in haystack or "crac" in haystack or "work order" in haystack:
        return "Experience / CRAC / work order document"
    if "bis" in haystack or "iso" in haystack or "type test" in haystack:
        return "BIS / ISO / type test certificate"
    if "biddoc" in haystack:
        return "Tender bid document"
    if "eligibility" in haystack or "atc" in haystack:
        return "ATC / eligibility criteria"
    return "Unclassified PDF"


def read_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    page_texts: list[str] = []
    for page in reader.pages:
        try:
            page_texts.append(page.extract_text() or "")
        except Exception:
            page_texts.append("")
    full_text = "\n".join(page_texts).strip()
    return {
        "page_count": len(reader.pages),
        "page_texts": page_texts,
        "text_chars": len(full_text),
        "pdf_type": "digital" if len(full_text) > 100 else "scanned",
        "text_extractable": len(full_text) > 100,
    }


def iter_pdfs() -> Iterable[Path]:
    return sorted(TENDER_ROOT.rglob("*.pdf"))


def normalize_id(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def format_status(status: str) -> str:
    return status


def find_ids(page_texts: list[str], regex: re.Pattern[str]) -> list[dict]:
    hits: list[dict] = []
    seen: set[tuple[str, int]] = set()
    for idx, text in enumerate(page_texts, start=1):
        for match in regex.finditer(text.upper()):
            value = normalize_id(match.group(0))
            key = (value, idx)
            if key not in seen:
                seen.add(key)
                hits.append({"value": value, "page": idx})
    return hits


def build_manifest() -> tuple[list[dict], dict[str, dict]]:
    rows: list[dict] = []
    pdf_cache: dict[str, dict] = {}
    for path in iter_pdfs():
        info = read_pdf(path)
        pdf_cache[path.as_posix()] = info
        bidder_id, bidder_name = bidder_from_path(path)
        rows.append(
            {
                "tender_id": TENDER_ID,
                "bidder_id": bidder_id,
                "bidder_name": bidder_name,
                "folder": path.parent.relative_to(TENDER_ROOT).as_posix(),
                "filename": path.name,
                "document_category": document_category(path),
                "expected_doc_type": expected_doc_type(path),
                "pdf_type": info["pdf_type"],
                "page_count": info["page_count"],
                "text_extractable": "Yes" if info["text_extractable"] else "No",
                "needs_ocr": "Yes" if info["pdf_type"] == "scanned" else "No",
                "notes": "Route to OCR" if info["pdf_type"] == "scanned" else "Use direct PDF text extraction",
            }
        )
    return rows, pdf_cache


def build_extractions(manifest_rows: list[dict], pdf_cache: dict[str, dict]) -> list[dict]:
    rows: list[dict] = []
    for item in manifest_rows:
        if item["document_category"] != "Bidder Submission":
            continue
        path = TENDER_ROOT / item["folder"] / item["filename"]
        info = pdf_cache[path.as_posix()]
        rel_file = f"{item['folder']}/{item['filename']}"
        doc_type = item["expected_doc_type"]
        if doc_type != "Unclassified PDF":
            rows.append(
                {
                    "tender_id": TENDER_ID,
                    "bidder_id": item["bidder_id"],
                    "bidder_name": item["bidder_name"],
                    "document_name": rel_file,
                    "field_name": "document_presence",
                    "extracted_value": doc_type,
                    "confidence": 0.85 if item["pdf_type"] == "scanned" else 0.95,
                    "page": 1 if item["page_count"] else "",
                    "source_method": "filename_inventory",
                    "validation_status": "present",
                    "notes": "Document type inferred from filename/text cue.",
                }
            )

        if not info["text_extractable"]:
            rows.append(
                {
                    "tender_id": TENDER_ID,
                    "bidder_id": item["bidder_id"],
                    "bidder_name": item["bidder_name"],
                    "document_name": rel_file,
                    "field_name": "ocr_required",
                    "extracted_value": "OCR required before field extraction",
                    "confidence": "",
                    "page": "",
                    "source_method": "pdf_classification",
                    "validation_status": "needs_review",
                    "notes": "No extractable text detected by pypdf.",
                }
            )
            continue

        for hit in find_ids(info["page_texts"], PAN_RE):
            rows.append(
                {
                    "tender_id": TENDER_ID,
                    "bidder_id": item["bidder_id"],
                    "bidder_name": item["bidder_name"],
                    "document_name": rel_file,
                    "field_name": "PAN",
                    "extracted_value": hit["value"],
                    "confidence": 0.95,
                    "page": hit["page"],
                    "source_method": "direct_pdf_text",
                    "validation_status": "format_valid",
                    "notes": "Exact PAN pattern match.",
                }
            )
        for hit in find_ids(info["page_texts"], GSTIN_RE):
            rows.append(
                {
                    "tender_id": TENDER_ID,
                    "bidder_id": item["bidder_id"],
                    "bidder_name": item["bidder_name"],
                    "document_name": rel_file,
                    "field_name": "GSTIN",
                    "extracted_value": hit["value"],
                    "confidence": 0.95,
                    "page": hit["page"],
                    "source_method": "direct_pdf_text",
                    "validation_status": "format_valid",
                    "notes": "Exact GSTIN pattern match; external API not used.",
                }
            )
    return rows


def first_by_type(manifest_rows: list[dict], bidder_id: str, doc_type: str) -> dict | None:
    for row in manifest_rows:
        if row["bidder_id"] == bidder_id and row["expected_doc_type"] == doc_type:
            return row
    return None


def first_field(extraction_rows: list[dict], bidder_id: str, field_name: str) -> dict | None:
    for row in extraction_rows:
        if row["bidder_id"] == bidder_id and row["field_name"] == field_name:
            return row
    return None


def field_candidates(extraction_rows: list[dict], bidder_id: str, field_name: str) -> list[dict]:
    return [row for row in extraction_rows if row["bidder_id"] == bidder_id and row["field_name"] == field_name]


def doc_priority(document_name: str) -> int:
    doc = document_name.lower()
    if "gst_registration" in doc or ("gst" in doc and "registration" in doc):
        return 0
    if "udyam" in doc or "integrity" in doc or "declaration" in doc:
        return 1
    if "mii" in doc or "localcontent" in doc or "local content" in doc:
        return 2
    if "experience" in doc or "crac" in doc:
        return 5
    return 3


def preferred_pan(extraction_rows: list[dict], bidder_id: str) -> dict | None:
    candidates = field_candidates(extraction_rows, bidder_id, "PAN")
    if not candidates:
        return None
    return sorted(candidates, key=lambda row: (doc_priority(row["document_name"]), row["page"] or 999))[0]


def preferred_gstin(extraction_rows: list[dict], bidder_id: str, pan_value: str | None = None) -> dict | None:
    candidates = field_candidates(extraction_rows, bidder_id, "GSTIN")
    if not candidates:
        return None
    if pan_value:
        matching = [row for row in candidates if row["extracted_value"][2:12] == pan_value]
        if matching:
            return sorted(matching, key=lambda row: (doc_priority(row["document_name"]), row["page"] or 999))[0]
    return sorted(candidates, key=lambda row: (doc_priority(row["document_name"]), row["page"] or 999))[0]


def bidder_ids(manifest_rows: list[dict]) -> list[tuple[str, str]]:
    seen: dict[str, str] = {}
    for row in manifest_rows:
        if row["bidder_id"]:
            seen[row["bidder_id"]] = row["bidder_name"]
    return sorted(seen.items())


def bidder_folder(manifest_rows: list[dict], bidder_id: str) -> str:
    for row in manifest_rows:
        if row["bidder_id"] == bidder_id:
            return row["folder"].split("/", 1)[0]
    return bidder_id


def compliance_for_requirement(
    req: Requirement,
    bidder_id: str,
    bidder_name: str,
    manifest_rows: list[dict],
    extraction_rows: list[dict],
) -> dict:
    if req.req_id == "T1-R01":
        pdf_count = sum(1 for row in manifest_rows if row["bidder_id"] == bidder_id)
        return {
            "status": "Present / Compliant" if pdf_count else "Needs Review - likely missing",
            "evidence_file": f"03_Bidder_Submissions/{bidder_folder(manifest_rows, bidder_id)}",
            "evidence_page": "N/A",
            "extracted_value": bidder_name,
            "notes": f"{pdf_count} submitted PDFs found.",
        }
    if req.req_id == "T1-R02":
        pan = preferred_pan(extraction_rows, bidder_id)
        return evidence_from_field(pan, "PAN not found in direct text; OCR/manual review required.")
    if req.req_id == "T1-R03":
        pan = preferred_pan(extraction_rows, bidder_id)
        gstin = preferred_gstin(extraction_rows, bidder_id, pan["extracted_value"] if pan else None)
        return evidence_from_field(gstin, "GSTIN not found in direct text; OCR/manual review required.")
    if req.req_id == "T1-R05":
        pan = preferred_pan(extraction_rows, bidder_id)
        gstin = preferred_gstin(extraction_rows, bidder_id, pan["extracted_value"] if pan else None)
        if pan and gstin:
            match = gstin["extracted_value"][2:12] == pan["extracted_value"]
            return {
                "status": "Present / Compliant" if match else "Needs Review - mismatch",
                "evidence_file": gstin["document_name"],
                "evidence_page": gstin["page"],
                "extracted_value": f"PAN={pan['extracted_value']}; GSTIN={gstin['extracted_value']}",
                "notes": "GSTIN chars 3-12 match PAN." if match else "GSTIN chars 3-12 do not match PAN.",
            }
        return {
            "status": "Needs Review - requires PAN and GSTIN",
            "evidence_file": "",
            "evidence_page": "",
            "extracted_value": "",
            "notes": "Both PAN and GSTIN are required for cross-check.",
        }

    doc_map = {
        "T1-R04": "GST registration certificate",
        "T1-R06": "Bidder turnover CA certificate",
        "T1-R07": "OEM authorization certificate",
        "T1-R08": "OEM turnover CA certificate",
        "T1-R09": "MII local content declaration",
        "T1-R10": "Integrity Pact or declaration document",
        "T1-R11": "Udyam / MSME certificate",
        "T1-R12": "BIS / ISO / type test certificate",
        "T1-R13": "Experience / CRAC / work order document",
    }
    doc_type = doc_map.get(req.req_id)
    if doc_type:
        found = first_by_type(manifest_rows, bidder_id, doc_type)
        if found:
            return {
                "status": "Present / Compliant",
                "evidence_file": f"{found['folder']}/{found['filename']}",
                "evidence_page": 1,
                "extracted_value": doc_type,
                "notes": "Document presence inferred from file inventory.",
            }
        return {
            "status": "Needs Review - likely missing",
            "evidence_file": "",
            "evidence_page": "",
            "extracted_value": "",
            "notes": "No matching document identified in bidder folder.",
        }
    return {
        "status": "Needs Review",
        "evidence_file": "",
        "evidence_page": "",
        "extracted_value": "",
        "notes": "No rule implemented.",
    }


def evidence_from_field(row: dict | None, missing_note: str) -> dict:
    if row:
        return {
            "status": "Present / Compliant",
            "evidence_file": row["document_name"],
            "evidence_page": row["page"],
            "extracted_value": row["extracted_value"],
            "notes": row["notes"],
        }
    return {
        "status": "Needs Review - likely missing",
        "evidence_file": "",
        "evidence_page": "",
        "extracted_value": "",
        "notes": missing_note,
    }


def build_compliance_rows(manifest_rows: list[dict], extraction_rows: list[dict]) -> list[dict]:
    rows: list[dict] = []
    bidders = bidder_ids(manifest_rows)
    for req in REQUIREMENTS:
        base = {
            "req_id": req.req_id,
            "requirement_text": req.requirement_text,
            "required_document": req.required_document,
            "validation_rule": req.validation_rule,
            "decision_rule": req.decision_rule,
            "remarks": req.remarks,
        }
        for bidder_id, bidder_name in bidders:
            result = compliance_for_requirement(req, bidder_id, bidder_name, manifest_rows, extraction_rows)
            prefix = bidder_id
            base[f"{prefix}_status"] = result["status"]
            base[f"{prefix}_evidence_file"] = result["evidence_file"]
            base[f"{prefix}_evidence_page"] = result["evidence_page"]
            base[f"{prefix}_extracted_value"] = result["extracted_value"]
            base[f"{prefix}_notes"] = result["notes"]
        rows.append(base)
    return rows


def autosize(ws) -> None:
    for column_cells in ws.columns:
        length = 0
        col = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            length = max(length, len(value))
        ws.column_dimensions[col].width = min(max(length + 2, 12), 55)


def write_workbook(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for title, (headers, rows) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        autosize(ws)
    wb.save(path)


def main() -> None:
    ensure_dirs()
    manifest_rows, pdf_cache = build_manifest()
    extraction_rows = build_extractions(manifest_rows, pdf_cache)
    compliance_rows = build_compliance_rows(manifest_rows, extraction_rows)

    manifest_headers = [
        "tender_id",
        "bidder_id",
        "bidder_name",
        "folder",
        "filename",
        "document_category",
        "expected_doc_type",
        "pdf_type",
        "page_count",
        "text_extractable",
        "needs_ocr",
        "notes",
    ]
    checklist_headers = [
        "req_id",
        "clause_ref",
        "requirement_text",
        "required_document",
        "field_to_extract",
        "validation_rule",
        "tier",
        "decision_rule",
        "remarks",
    ]
    ground_truth_headers = [
        "tender_id",
        "bidder_id",
        "field_name",
        "expected_value",
        "document_name",
        "evidence_page",
        "evidence_text_or_region",
        "verified_by",
        "verified_date",
        "notes",
    ]
    gt_rows = [
        {
            "tender_id": TENDER_ID,
            "bidder_id": bidder_id,
            "field_name": req.field_to_extract,
            "expected_value": "",
            "document_name": "",
            "evidence_page": "",
            "evidence_text_or_region": "",
            "verified_by": "",
            "verified_date": "",
            "notes": f"{req.req_id}: {req.requirement_text}",
        }
        for bidder_id, _bidder_name in bidder_ids(manifest_rows)
        for req in REQUIREMENTS
    ]
    checklist_rows = [req.__dict__ for req in REQUIREMENTS]

    write_workbook(CONTROL_DIR / "pilot_manifest.xlsx", {"Manifest": (manifest_headers, manifest_rows)})
    write_workbook(
        CONTROL_DIR / "tier1_requirement_checklist.xlsx",
        {"Tier1 Checklist": (checklist_headers, checklist_rows)},
    )
    write_workbook(
        CONTROL_DIR / "ground_truth_tier1.xlsx",
        {"Ground Truth": (ground_truth_headers, gt_rows)},
    )

    extraction_headers = [
        "tender_id",
        "bidder_id",
        "bidder_name",
        "document_name",
        "field_name",
        "extracted_value",
        "confidence",
        "page",
        "source_method",
        "validation_status",
        "notes",
    ]
    write_workbook(
        EXTRACTION_DIR / "extracted_candidates.xlsx",
        {"Extracted Candidates": (extraction_headers, extraction_rows)},
    )
    (EXTRACTION_DIR / "extracted_candidates.json").write_text(
        json.dumps(extraction_rows, indent=2),
        encoding="utf-8",
    )

    compliance_headers = list(compliance_rows[0].keys()) if compliance_rows else []
    write_workbook(
        MATRIX_DIR / "tier1_compliance_matrix.xlsx",
        {"Tier1 Matrix": (compliance_headers, compliance_rows)},
    )

    audit_headers = [
        "timestamp",
        "tender_id",
        "bidder_id",
        "requirement_id",
        "extracted_value",
        "confidence",
        "evidence_page",
        "system_decision",
        "decision_basis",
        "reviewer_id",
        "human_decision",
        "override_flag",
        "override_reason",
    ]
    write_workbook(AUDIT_DIR / "audit_log_template.xlsx", {"Audit Log": (audit_headers, [])})

    print(f"Generated {len(manifest_rows)} manifest rows")
    print(f"Generated {len(extraction_rows)} extraction candidate rows")
    print(f"Generated {len(compliance_rows)} compliance rows")
    print(f"Output root: {TENDER_ROOT}")


if __name__ == "__main__":
    main()
