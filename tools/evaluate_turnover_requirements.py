from __future__ import annotations

import argparse
import json
import os
import re
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TENDER_ID = "Tender_01_VacuumFlask"

STATUS_LIKELY_MEETS = "Likely Meets \u2014 Needs Human Review"
STATUS_LIKELY_NOT = "Likely Does Not Meet \u2014 Needs Human Review"
STATUS_UNCLEAR = "Unclear \u2014 Needs Human Review"
STATUS_NOT_FOUND = "Evidence Not Found \u2014 Needs Human Review"

TURNOVER_DIR_NAME = "07_turnover_evaluation"
DEFAULT_GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"
DEFAULT_GROQ_MODEL = "llama-3.3-70b-versatile"

REQUIREMENT_HEADERS = [
    "tender_id",
    "requirement_type",
    "requirement_label",
    "threshold_text",
    "threshold_amount",
    "threshold_unit",
    "threshold_period_years",
    "currency",
    "source_document_id",
    "source_file",
    "source_page",
    "source_text_snippet",
    "extraction_method",
    "confidence",
    "notes",
]

EVIDENCE_HEADERS = [
    "evidence_id",
    "tender_id",
    "bidder_id",
    "bidder_name",
    "requirement_type",
    "document_id",
    "document_name",
    "document_type_guess",
    "source_page",
    "evidence_text",
    "amount_candidates",
    "source_method",
    "evidence_score",
    "notes",
]

LLM_HEADERS = [
    "tender_id",
    "bidder_id",
    "bidder_name",
    "requirement_type",
    "threshold_amount",
    "threshold_period_years",
    "currency",
    "extracted_values",
    "computed_average_amount",
    "meets_requirement_suggestion",
    "review_status",
    "confidence",
    "reasoning_summary",
    "human_review_required",
    "source_file",
    "source_page",
    "source_quote",
    "provider_used",
    "api_attempted",
    "api_error",
]

MATRIX_HEADERS = [
    "tender_id",
    "bidder_id",
    "bidder_name",
    "bidder_turnover_threshold",
    "bidder_turnover_extracted_average",
    "bidder_turnover_status",
    "bidder_turnover_source_file",
    "bidder_turnover_source_page",
    "bidder_turnover_evidence_text",
    "oem_turnover_threshold",
    "oem_turnover_extracted_average",
    "oem_turnover_status",
    "oem_turnover_source_file",
    "oem_turnover_source_page",
    "oem_turnover_evidence_text",
    "llm_confidence",
    "human_review_required",
    "reviewer_decision",
    "reviewer_notes",
]


def resolve_tender_root(tender_id: str, tender_root: str | Path | None = None) -> Path:
    return Path(tender_root).resolve() if tender_root else ROOT / tender_id


def paths(tender_id: str, tender_root: str | Path | None = None) -> dict[str, Path]:
    tender_root = resolve_tender_root(tender_id, tender_root)
    output_root = tender_root / "05_Extraction_Output"
    turnover_root = output_root / TURNOVER_DIR_NAME
    return {
        "manifest": output_root / "document_manifest.json",
        "text_index": output_root / "01_text_extraction" / "all_document_text_index.json",
        "requirements": output_root / "05_tender_requirements" / "required_document_attributes.json",
        "turnover_root": turnover_root,
        "tender_requirements_xlsx": turnover_root / "tender_turnover_requirements.xlsx",
        "tender_requirements_json": turnover_root / "tender_turnover_requirements.json",
        "evidence_xlsx": turnover_root / "bidder_turnover_evidence.xlsx",
        "evidence_json": turnover_root / "bidder_turnover_evidence.json",
        "llm_results_xlsx": turnover_root / "turnover_llm_results.xlsx",
        "llm_results_json": turnover_root / "turnover_llm_results.json",
        "matrix_xlsx": turnover_root / "turnover_review_matrix.xlsx",
        "matrix_json": turnover_root / "turnover_review_matrix.json",
        "notes_json": turnover_root / "turnover_notes.json",
        "api_log": turnover_root / "llm_api_log.jsonl",
    }


def load_env_file(env_path: Path | None = None) -> None:
    env_path = env_path or ROOT / "tools" / ".env"
    if not env_path.exists():
        return
    try:
        from dotenv import load_dotenv  # type: ignore

        load_dotenv(env_path, override=False)
        return
    except Exception:
        pass

    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")


def excel_safe(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", value)
    if isinstance(value, (dict, list)):
        return ILLEGAL_CHARACTERS_RE.sub(" ", json.dumps(value, ensure_ascii=False))
    return value


def autosize(ws) -> None:
    for column_cells in ws.columns:
        column = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[column].width = min(max(max_len + 2, 12), 70)


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append([excel_safe(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)
    wb.save(path)


def normalize_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def snippet(text: str, start: int, end: int, width: int = 320) -> str:
    left = max(0, start - width)
    right = min(len(text), end + width)
    return normalize_spaces(text[left:right])


def normalize_amount(amount_text: str) -> tuple[float | None, str]:
    text = normalize_spaces(amount_text).casefold()
    match = re.search(r"([0-9][0-9,]*(?:\.[0-9]+)?)\s*(lakh|lakhs|lac|lacs|crore|crores|cr|rs|inr)?", text)
    if not match:
        return None, ""
    number = float(match.group(1).replace(",", ""))
    unit = match.group(2) or "INR"
    multiplier = 1
    if unit in {"lakh", "lakhs", "lac", "lacs"}:
        multiplier = 100000
        unit = "Lakh"
    elif unit in {"crore", "crores", "cr"}:
        multiplier = 10000000
        unit = "Crore"
    else:
        unit = "INR"
    return number * multiplier, unit


def looks_like_turnover_amount(raw_text: str, number_text: str, context: str) -> bool:
    raw_lower = raw_text.casefold()
    context_lower = context.casefold()
    if re.search(r"rs\.?|inr|\u20b9|lakh|lakhs|lac|lacs|crore|crores|cr", raw_lower):
        return True

    number_clean = number_text.strip(",")
    indian_grouped_number = bool(re.fullmatch(r"\d{1,3}(?:,\d{2})*,\d{3}", number_clean))
    turnover_context = bool(
        re.search(
            r"turnover|amount|financial|gross|annual|sales|revenue|certificate|fy|year",
            context_lower,
        )
    )
    contact_context = bool(re.search(r"\+91|mob\.?|mobile|phone|tel\.?|email|e-mail", context_lower))
    return indian_grouped_number and turnover_context and not contact_context


def amount_candidates(text: str) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    amount_pattern = re.compile(
        r"(?:rs\.?|inr|₹)?\s*([0-9][0-9,]*(?:\.[0-9]+)?)\s*(lakh|lakhs|lac|lacs|crore|crores|cr)?",
        flags=re.IGNORECASE,
    )
    for match in amount_pattern.finditer(text):
        raw = match.group(0).strip()
        context = snippet(text, match.start(), match.end(), width=90)
        if not looks_like_turnover_amount(raw, match.group(1), context):
            continue
        amount, unit = normalize_amount(raw)
        if amount is None:
            continue
        # Avoid obvious dates, years, phone fragments, and tiny unrelated numbers.
        if 1900 <= amount <= 2100 and not match.group(2):
            continue
        if amount < 1000 and not match.group(2):
            continue
        candidates.append(
            {
                "raw_text": raw,
                "amount": amount,
                "unit": unit,
                "context": context,
            }
        )
    return candidates


def extract_financial_year(text: str) -> str:
    patterns = [
        r"(20[0-9]{2}\s*-\s*[0-9]{2})",
        r"(FY\s*20[0-9]{2}\s*-\s*[0-9]{2})",
        r"year\s+([0-9]{2}\s*-\s*[0-9]{2})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return normalize_spaces(match.group(1).upper())
    return ""


def manifest_lookup(manifest_rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {row["document_id"]: row for row in manifest_rows}


def extract_turnover_requirements(tender_id: str, text_index_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    patterns = [
        (
            "bidder_turnover",
            "Bidder Minimum Average Annual Turnover",
            r"Minimum\s+Average\s+Annual\s+Turnover\s+of\s+the\s+bidder\s*\(For\s+3\s+Years\)\s*(?P<threshold>[0-9][0-9,.\s]*(?:Lakh|Lakhs|Lac|Crore|Crores|Cr|INR|Rs)?\s*(?:\([sS]\))?)",
        ),
        (
            "oem_turnover",
            "OEM Average Turnover",
            r"OEM\s+Average\s+Turnover\s*\(Last\s+3\s+Years\)\s*(?P<threshold>[0-9][0-9,.\s]*(?:Lakh|Lakhs|Lac|Crore|Crores|Cr|INR|Rs)?\s*(?:\([sS]\))?)",
        ),
    ]
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for page in text_index_rows:
        if page.get("document_side") != "tender-side":
            continue
        text = page.get("chosen_text", "") or ""
        for requirement_type, label, pattern in patterns:
            if requirement_type in seen:
                continue
            match = re.search(pattern, text, flags=re.IGNORECASE)
            if not match:
                continue
            threshold_text = normalize_spaces(match.group("threshold"))
            threshold_amount, threshold_unit = normalize_amount(threshold_text)
            rows.append(
                {
                    "tender_id": tender_id,
                    "requirement_type": requirement_type,
                    "requirement_label": label,
                    "threshold_text": threshold_text,
                    "threshold_amount": threshold_amount,
                    "threshold_unit": threshold_unit,
                    "threshold_period_years": 3,
                    "currency": "INR",
                    "source_document_id": page.get("document_id", ""),
                    "source_file": page.get("file_name", ""),
                    "source_page": page.get("page_number", ""),
                    "source_text_snippet": snippet(text, match.start(), match.end()),
                    "extraction_method": "tender_label_regex",
                    "confidence": 0.95,
                    "notes": "",
                }
            )
            seen.add(requirement_type)
    return rows


def is_turnover_doc(doc: dict[str, Any], requirement_type: str) -> bool:
    haystack = normalize_spaces(
        f"{doc.get('document_type_guess', '')} {doc.get('file_name', '')} {doc.get('folder_path', '')}"
    ).casefold()
    if doc.get("document_side") != "bidder-side":
        return False
    if requirement_type == "bidder_turnover":
        return "turnover" in haystack and "oem" not in haystack
    if requirement_type == "oem_turnover":
        return "oem" in haystack and "turnover" in haystack
    return False


def relevant_turnover_text(text: str) -> bool:
    return bool(re.search(r"turnover|financial|audited|balance|profit|loss|chartered|certificate|lakh|crore|revenue|sales", text, re.IGNORECASE))


def build_evidence_rows(
    tender_id: str,
    manifest_rows: list[dict[str, Any]],
    text_index_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    docs_by_id = manifest_lookup(manifest_rows)
    rows: list[dict[str, Any]] = []
    counter = 1
    for page in text_index_rows:
        doc = docs_by_id.get(page.get("document_id", ""), {})
        if doc.get("document_side") != "bidder-side":
            continue
        text = page.get("chosen_text", "") or ""
        if not relevant_turnover_text(text):
            continue
        for requirement_type in ("bidder_turnover", "oem_turnover"):
            if not is_turnover_doc(doc, requirement_type):
                continue
            amounts = amount_candidates(text)
            evidence_text = compact_evidence_text(text)
            rows.append(
                {
                    "evidence_id": f"TURN-EVID-{counter:04d}",
                    "tender_id": tender_id,
                    "bidder_id": doc.get("bidder_id", ""),
                    "bidder_name": doc.get("bidder_name", ""),
                    "requirement_type": requirement_type,
                    "document_id": doc.get("document_id", ""),
                    "document_name": doc.get("file_name", ""),
                    "document_type_guess": doc.get("document_type_guess", ""),
                    "source_page": page.get("page_number", ""),
                    "evidence_text": evidence_text,
                    "amount_candidates": json.dumps(amounts, ensure_ascii=False),
                    "source_method": page.get("chosen_text_source", ""),
                    "evidence_score": evidence_score(text, amounts),
                    "notes": "",
                }
            )
            counter += 1
    return rows


def compact_evidence_text(text: str, max_chars: int = 2500) -> str:
    lines = [normalize_spaces(line) for line in text.splitlines()]
    lines = [line for line in lines if line and relevant_turnover_text(line)]
    if not lines:
        compact = normalize_spaces(text)
    else:
        compact = " | ".join(lines)
    return compact[:max_chars]


def evidence_score(text: str, amounts: list[dict[str, Any]]) -> float:
    score = 0.3
    if re.search(r"turnover", text, re.IGNORECASE):
        score += 0.2
    if re.search(r"chartered|ca |certificate", text, re.IGNORECASE):
        score += 0.15
    if amounts:
        score += 0.25
    return round(min(score, 0.95), 2)


def redact_sensitive(text: str) -> str:
    pan_re = re.compile(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b")
    gstin_re = re.compile(r"\b[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b")
    email_re = re.compile(r"\b[\w.\-]+@[\w.\-]+\.\w+\b")
    phone_re = re.compile(r"(?<!\d)(?:\+91[-\s]?)?[6-9]\d{9}(?!\d)")
    text = gstin_re.sub("[REDACTED_GSTIN]", text)
    text = pan_re.sub("[REDACTED_PAN]", text)
    text = email_re.sub("[REDACTED_EMAIL]", text)
    text = phone_re.sub("[REDACTED_PHONE]", text)
    return text


def llm_payload(
    tender_id: str,
    bidder_id: str,
    bidder_name: str,
    requirement: dict[str, Any],
    evidence_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "tender_id": tender_id,
        "bidder_id": bidder_id,
        "bidder_name": bidder_name,
        "requirement_type": requirement["requirement_type"],
        "threshold_amount": requirement["threshold_amount"],
        "threshold_period_years": requirement["threshold_period_years"],
        "currency": "INR",
        "evidence": [
            {
                "document_name": row["document_name"],
                "document_type_guess": row["document_type_guess"],
                "source_page": row["source_page"],
                "evidence_text": redact_sensitive(row["evidence_text"]),
                "amount_candidates": row["amount_candidates"],
            }
            for row in evidence_rows
        ],
        "instructions": [
            "Return strict JSON only.",
            "Do not infer values from weak wording.",
            "Do not treat CA firm turnover as bidder turnover.",
            "Do not treat bidder turnover as OEM turnover.",
            "Always cite source page and quote.",
        ],
    }


def openai_compatible_llm_call(payload: dict[str, Any], api_log: Path) -> tuple[dict[str, Any] | None, str]:
    groq_key = os.environ.get("GROQ_API_KEY")
    api_url = os.environ.get("TURNOVER_LLM_API_URL")
    if not api_url:
        api_url = DEFAULT_GROQ_API_URL if groq_key else "https://api.openai.com/v1/chat/completions"
    api_key = os.environ.get("TURNOVER_LLM_API_KEY") or groq_key or os.environ.get("OPENAI_API_KEY")
    model = os.environ.get("TURNOVER_LLM_MODEL") or (DEFAULT_GROQ_MODEL if groq_key else "gpt-4.1-mini")
    if not api_key:
        return None, "No API key found in TURNOVER_LLM_API_KEY, GROQ_API_KEY, or OPENAI_API_KEY."

    system_prompt = (
        "You extract turnover values from procurement evidence. "
        "Return only JSON matching the requested contract. "
        "All conclusions require human review."
    )
    request_body = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
        ],
        "response_format": {"type": "json_object"},
        "temperature": 0,
    }
    api_log.parent.mkdir(parents=True, exist_ok=True)
    with api_log.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps({"ts": time.time(), "type": "request", "payload": request_body}, ensure_ascii=False) + "\n")
    request = urllib.request.Request(
        api_url,
        data=json.dumps(request_body).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            raw = response.read().decode("utf-8", errors="replace")
    except (urllib.error.URLError, TimeoutError, ValueError) as exc:
        return None, str(exc)
    with api_log.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps({"ts": time.time(), "type": "response", "raw": raw}, ensure_ascii=False) + "\n")
    try:
        response_json = json.loads(raw)
        content = response_json["choices"][0]["message"]["content"]
        return json.loads(content), ""
    except Exception as exc:
        return None, f"Could not parse LLM response: {exc}"


def offline_turnover_result(
    requirement: dict[str, Any],
    evidence_rows: list[dict[str, Any]],
    api_error: str = "",
) -> dict[str, Any]:
    threshold = float(requirement.get("threshold_amount") or 0)
    period_years = int(requirement.get("threshold_period_years") or 3)
    extracted_values: list[dict[str, Any]] = []
    source_file = ""
    source_page = ""
    source_quote = ""

    for evidence in sorted(evidence_rows, key=lambda row: row.get("evidence_score", 0), reverse=True):
        source_file = source_file or evidence.get("document_name", "")
        source_page = source_page or evidence.get("source_page", "")
        source_quote = source_quote or evidence.get("evidence_text", "")
        try:
            amounts = json.loads(evidence.get("amount_candidates", "[]"))
        except json.JSONDecodeError:
            amounts = []
        for amount in amounts:
            extracted_values.append(
                {
                    "financial_year": extract_financial_year(amount.get("context", "")),
                    "amount": amount.get("amount"),
                    "unit": "INR",
                    "source_quote": amount.get("context", ""),
                    "source_page": evidence.get("source_page", ""),
                }
            )

    # Keep the offline fallback conservative: it computes only when at least one explicit amount is found.
    numeric_values = [float(row["amount"]) for row in extracted_values if row.get("amount")]
    if not evidence_rows:
        suggestion = "evidence_not_found"
        review_status = STATUS_NOT_FOUND
        confidence = 0.0
        summary = "No likely turnover evidence document was identified."
        average = None
    elif not numeric_values:
        suggestion = "unclear"
        review_status = STATUS_UNCLEAR
        confidence = 0.25
        summary = "Likely turnover evidence exists, but explicit numeric values were not reliably extracted."
        average = None
    elif len(numeric_values) < period_years and not re.search(r"\b(avg|average)\b|औसत", source_quote, flags=re.IGNORECASE):
        suggestion = "unclear"
        review_status = STATUS_UNCLEAR
        confidence = 0.35
        summary = (
            f"Only {len(numeric_values)} explicit turnover amount(s) were extracted for a "
            f"{period_years}-year average requirement; human review is required."
        )
        average = None
    else:
        average = sum(numeric_values[:period_years]) / min(len(numeric_values), period_years)
        if average >= threshold:
            suggestion = "likely_meets"
            review_status = STATUS_LIKELY_MEETS
        else:
            suggestion = "likely_does_not_meet"
            review_status = STATUS_LIKELY_NOT
        confidence = 0.45 if len(numeric_values) < 3 else 0.6
        summary = "Offline fallback computed a tentative average from explicit amount candidates; human review is mandatory."

    return {
        "requirement_type": requirement["requirement_type"],
        "threshold_amount": threshold,
        "threshold_period_years": requirement["threshold_period_years"],
        "currency": "INR",
        "extracted_values": extracted_values,
        "computed_average_amount": average,
        "meets_requirement_suggestion": suggestion,
        "review_status": review_status,
        "confidence": confidence,
        "reasoning_summary": summary,
        "human_review_required": True,
        "source_file": source_file,
        "source_page": source_page,
        "source_quote": source_quote,
        "provider_used": "offline_fallback",
        "api_attempted": False,
        "api_error": api_error,
    }


def normalize_llm_result(
    requirement: dict[str, Any],
    llm_result: dict[str, Any],
    evidence_rows: list[dict[str, Any]],
    provider_used: str,
    api_error: str,
) -> dict[str, Any]:
    suggestion = llm_result.get("meets_requirement_suggestion", "unclear")
    status_map = {
        "likely_meets": STATUS_LIKELY_MEETS,
        "likely_does_not_meet": STATUS_LIKELY_NOT,
        "unclear": STATUS_UNCLEAR,
        "evidence_not_found": STATUS_NOT_FOUND,
    }
    top_evidence = evidence_rows[0] if evidence_rows else {}
    return {
        "requirement_type": llm_result.get("requirement_type", requirement["requirement_type"]),
        "threshold_amount": llm_result.get("threshold_amount", requirement["threshold_amount"]),
        "threshold_period_years": llm_result.get("threshold_period_years", requirement["threshold_period_years"]),
        "currency": llm_result.get("currency", "INR"),
        "extracted_values": llm_result.get("extracted_values", []),
        "computed_average_amount": llm_result.get("computed_average_amount"),
        "meets_requirement_suggestion": suggestion,
        "review_status": status_map.get(suggestion, STATUS_UNCLEAR),
        "confidence": llm_result.get("confidence", 0),
        "reasoning_summary": llm_result.get("reasoning_summary", ""),
        "human_review_required": True,
        "source_file": top_evidence.get("document_name", ""),
        "source_page": top_evidence.get("source_page", ""),
        "source_quote": top_evidence.get("evidence_text", ""),
        "provider_used": provider_used,
        "api_attempted": True,
        "api_error": api_error,
    }


def evaluate_with_llm_or_fallback(
    tender_id: str,
    bidders: list[tuple[str, str]],
    requirements: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    api_log: Path,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    provider = os.environ.get("TURNOVER_LLM_PROVIDER", "openai_compatible")
    for bidder_id, bidder_name in bidders:
        for requirement in requirements:
            relevant = [
                row
                for row in evidence_rows
                if row["bidder_id"] == bidder_id and row["requirement_type"] == requirement["requirement_type"]
            ]
            relevant = sorted(relevant, key=lambda row: row.get("evidence_score", 0), reverse=True)[:4]
            payload = llm_payload(tender_id, bidder_id, bidder_name, requirement, relevant)
            llm_json: dict[str, Any] | None = None
            api_error = ""
            if provider == "openai_compatible":
                llm_json, api_error = openai_compatible_llm_call(payload, api_log)
            if llm_json:
                normalized = normalize_llm_result(requirement, llm_json, relevant, provider, api_error)
            else:
                normalized = offline_turnover_result(requirement, relevant, api_error=api_error)
            results.append(
                {
                    "tender_id": tender_id,
                    "bidder_id": bidder_id,
                    "bidder_name": bidder_name,
                    **normalized,
                }
            )
    return results


def bidders_from_manifest(manifest_rows: list[dict[str, Any]]) -> list[tuple[str, str]]:
    seen: dict[str, str] = {}
    for row in manifest_rows:
        bidder_id = row.get("bidder_id", "")
        if bidder_id:
            seen[bidder_id] = row.get("bidder_name", "")
    return sorted(seen.items())


def requirement_by_type(requirements: list[dict[str, Any]], requirement_type: str) -> dict[str, Any]:
    for requirement in requirements:
        if requirement["requirement_type"] == requirement_type:
            return requirement
    return {}


def result_by_type(results: list[dict[str, Any]], bidder_id: str, requirement_type: str) -> dict[str, Any]:
    for result in results:
        if result["bidder_id"] == bidder_id and result["requirement_type"] == requirement_type:
            return result
    return {}


def build_review_matrix(
    tender_id: str,
    bidders: list[tuple[str, str]],
    requirements: list[dict[str, Any]],
    llm_results: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    bidder_req = requirement_by_type(requirements, "bidder_turnover")
    oem_req = requirement_by_type(requirements, "oem_turnover")
    rows: list[dict[str, Any]] = []
    for bidder_id, bidder_name in bidders:
        bidder_result = result_by_type(llm_results, bidder_id, "bidder_turnover")
        oem_result = result_by_type(llm_results, bidder_id, "oem_turnover")
        confidences = [
            float(result.get("confidence", 0))
            for result in (bidder_result, oem_result)
            if result
        ]
        rows.append(
            {
                "tender_id": tender_id,
                "bidder_id": bidder_id,
                "bidder_name": bidder_name,
                "bidder_turnover_threshold": bidder_req.get("threshold_text", ""),
                "bidder_turnover_extracted_average": bidder_result.get("computed_average_amount", ""),
                "bidder_turnover_status": bidder_result.get("review_status", STATUS_NOT_FOUND),
                "bidder_turnover_source_file": bidder_result.get("source_file", ""),
                "bidder_turnover_source_page": bidder_result.get("source_page", ""),
                "bidder_turnover_evidence_text": bidder_result.get("source_quote", ""),
                "oem_turnover_threshold": oem_req.get("threshold_text", ""),
                "oem_turnover_extracted_average": oem_result.get("computed_average_amount", ""),
                "oem_turnover_status": oem_result.get("review_status", STATUS_NOT_FOUND),
                "oem_turnover_source_file": oem_result.get("source_file", ""),
                "oem_turnover_source_page": oem_result.get("source_page", ""),
                "oem_turnover_evidence_text": oem_result.get("source_quote", ""),
                "llm_confidence": round(sum(confidences) / len(confidences), 2) if confidences else 0,
                "human_review_required": "Yes",
                "reviewer_decision": "",
                "reviewer_notes": "",
            }
        )
    return rows


def validate_outputs(
    requirements: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    llm_results: list[dict[str, Any]],
    review_rows: list[dict[str, Any]],
    bidders: list[tuple[str, str]],
) -> list[str]:
    issues: list[str] = []
    found_types = {row["requirement_type"] for row in requirements}
    for requirement_type in ("bidder_turnover", "oem_turnover"):
        if requirement_type not in found_types:
            issues.append(f"Missing tender turnover requirement: {requirement_type}")
    if len(review_rows) != len(bidders):
        issues.append("Turnover review matrix row count does not match bidder count.")
    result_pairs = {(row["bidder_id"], row["requirement_type"]) for row in llm_results}
    for bidder_id, _name in bidders:
        for requirement_type in ("bidder_turnover", "oem_turnover"):
            if (bidder_id, requirement_type) not in result_pairs:
                issues.append(f"Missing LLM/fallback result for {bidder_id} {requirement_type}.")
    forbidden = ("rejected", "accepted", "disqualified", "non-compliant")
    for row in review_rows:
        text = json.dumps(row, ensure_ascii=False).casefold()
        for term in forbidden:
            if term in text:
                issues.append(f"Forbidden final decision wording '{term}' in review matrix.")
    return issues


def notes(
    tender_id: str,
    requirements: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    llm_results: list[dict[str, Any]],
    review_rows: list[dict[str, Any]],
    issues: list[str],
) -> dict[str, Any]:
    providers = sorted({row.get("provider_used", "") for row in llm_results})
    return {
        "tender_id": tender_id,
        "turnover_requirements": len(requirements),
        "evidence_rows": len(evidence_rows),
        "llm_result_rows": len(llm_results),
        "review_matrix_rows": len(review_rows),
        "providers_used": providers,
        "api_attempted": any(row.get("api_attempted") for row in llm_results),
        "validation_issues": issues,
        "notes": [
            "Only turnover is evaluated in this step.",
            "All turnover results require human review.",
            "If no API key is configured, offline fallback results are generated conservatively.",
            "Snippets are redacted before API calls.",
        ],
    }


def run(tender_id: str, tender_root: str | Path | None = None) -> int:
    load_env_file()
    path_map = paths(tender_id, tender_root)
    path_map["turnover_root"].mkdir(parents=True, exist_ok=True)
    manifest_rows = load_json(path_map["manifest"])
    text_index_rows = load_json(path_map["text_index"])
    _required_document_attrs = load_json(path_map["requirements"])

    requirements = extract_turnover_requirements(tender_id, text_index_rows)
    evidence_rows = build_evidence_rows(tender_id, manifest_rows, text_index_rows)
    bidders = bidders_from_manifest(manifest_rows)
    llm_results = evaluate_with_llm_or_fallback(
        tender_id,
        bidders,
        requirements,
        evidence_rows,
        path_map["api_log"],
    )
    review_rows = build_review_matrix(tender_id, bidders, requirements, llm_results)
    issues = validate_outputs(requirements, evidence_rows, llm_results, review_rows, bidders)
    notes_payload = notes(tender_id, requirements, evidence_rows, llm_results, review_rows, issues)

    write_xlsx(path_map["tender_requirements_xlsx"], "Tender Turnover Requirements", requirements, REQUIREMENT_HEADERS)
    write_json(path_map["tender_requirements_json"], requirements)
    write_xlsx(path_map["evidence_xlsx"], "Bidder Turnover Evidence", evidence_rows, EVIDENCE_HEADERS)
    write_json(path_map["evidence_json"], evidence_rows)
    write_xlsx(path_map["llm_results_xlsx"], "Turnover LLM Results", llm_results, LLM_HEADERS)
    write_json(path_map["llm_results_json"], llm_results)
    write_xlsx(path_map["matrix_xlsx"], "Turnover Review Matrix", review_rows, MATRIX_HEADERS)
    write_json(path_map["matrix_json"], review_rows)
    write_json(path_map["notes_json"], notes_payload)

    print(f"Tender turnover requirements: {len(requirements)}")
    print(f"Evidence rows: {len(evidence_rows)}")
    print(f"LLM/fallback result rows: {len(llm_results)}")
    print(f"Review matrix rows: {len(review_rows)}")
    print(f"Validation issues: {len(issues)}")
    for issue in issues:
        print(f"- {issue}")
    return 1 if issues else 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate bidder and OEM turnover requirements.")
    parser.add_argument("--tender-id", default=DEFAULT_TENDER_ID)
    parser.add_argument("--tender-root", default="", help="Optional explicit tender workspace root.")
    args = parser.parse_args()
    raise SystemExit(run(args.tender_id, args.tender_root or None))


if __name__ == "__main__":
    main()
